# -*- coding: utf-8 -*-
"""Genera DCIP_2.3_Arquitectura_Maestra_v3.0.xlsx a partir de dcip_model.py."""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.comments import Comment

import dcip_model as M

OUT = "/home/user/defensa/DCIP_2.3/DCIP_2.3_Arquitectura_Maestra_v3.0.xlsx"

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True, color=M.NAVY)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
F_NOTE = Font(name=FONT, size=9, italic=True, color="616161")
FILL_HDR = PatternFill("solid", fgColor=M.NAVY)
FILL_INPUT = PatternFill("solid", fgColor="FFFFCC")
FILL_EX = PatternFill("solid", fgColor="F5F5F5")
FILL_GREY = PatternFill("solid", fgColor="ECEFF1")
THIN = Side(style="thin", color="B0BEC5")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()
wb.remove(wb.active)


def sheet(name, widths=None, freeze="A2"):
    ws = wb.create_sheet(name)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    if freeze:
        ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    return ws


def header(ws, row, labels, fill=FILL_HDR):
    for c, lab in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=lab)
        cell.font = F_HDR
        cell.fill = fill
        cell.alignment = CENTER
        cell.border = BORDER


def put(ws, row, col, value, font=F_BASE, fill=None, align=WRAP, border=True, fmt=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font
    if fill:
        cell.fill = fill
    cell.alignment = align
    if border:
        cell.border = BORDER
    if fmt:
        cell.number_format = fmt
    return cell


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_NOTE


def dv_list(ws, formula, rng, allow_blank=True):
    dv = DataValidation(type="list", formula1=formula, allow_blank=allow_blank)
    dv.error = "Valor no permitido. Use la lista desplegable."
    dv.errorTitle = "DCIP 2.3"
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv


def dv_num(ws, rng, lo=0, hi=100):
    dv = DataValidation(type="decimal", operator="between", formula1=str(lo), formula2=str(hi), allow_blank=True)
    dv.error = f"La puntuación debe estar entre {lo} y {hi}."
    dv.errorTitle = "DCIP 2.3"
    ws.add_data_validation(dv)
    dv.add(rng)


def dim_fill(dim):
    return PatternFill("solid", fgColor=M.DIMENSIONES[dim]["color_suave"])


N_IND = len(M.INDICADORES)
LAST = N_IND + 1  # última fila de indicadores (encabezado en 1)

# =====================================================================================
# 16_LISTAS (se crea primero en memoria para poder referenciarla; se mueve al final)
# =====================================================================================
L = sheet("16_LISTAS", [16, 4, 4, 12, 10, 4, 10, 16, 10, 4, 16, 4, 14, 4, 14, 4, 20, 4, 22, 4, 16], freeze=None)
title(L, "16_LISTAS — Listas de validación y tablas de consulta", "No modificar sin actualizar la metodología. Las fórmulas de otras hojas dependen de estas posiciones.")
# A: confianza (orden ascendente 1..5)
header(L, 3, ["CONFIANZA", "", "", "NIVEL_FUENTE", "CONF_MAX_NUM", "", "BANDA_MIN", "BANDA", "BANDA_MAX", "", "TIPO_EVIDENCIA", "", "PROB_IMPACTO", "", "ESTADO_CONTROL", "", "CALIFICACION_CAMBIO", "", "ESCENARIO / HORIZONTE", "", "S_N"])
for i, n in enumerate(M.NIVELES_CONFIANZA):
    put(L, 4 + i, 1, n)
for i, (cod, nom, desc, cmax) in enumerate(M.NIVELES_FUENTE):
    put(L, 4 + i, 4, cod)
    put(L, 4 + i, 5, M.CONF_NUM[cmax])
for i, (lo, hi, nom, _, _) in enumerate(M.BANDAS):
    put(L, 4 + i, 7, lo)
    put(L, 4 + i, 8, nom)
    put(L, 4 + i, 9, hi)
for i, t in enumerate(M.TIPOS_EVIDENCIA):
    put(L, 4 + i, 11, t)
for i, t in enumerate(["BAJA", "MEDIA", "ALTA"]):
    put(L, 4 + i, 13, t)
for i, t in enumerate(["PENDIENTE", "CUMPLE", "NO CUMPLE", "NO APLICA"]):
    put(L, 4 + i, 15, t)
for i, t in enumerate(["MEJORA", "ESTABLE", "DETERIORO", "DETERIORO CRÍTICO", "NUEVA INFORMACIÓN"]):
    put(L, 4 + i, 17, t)
for i, t in enumerate(["BASE", "ALCISTA", "BAJISTA", "12–24 meses", "3–5 años", "5–10 años"]):
    put(L, 4 + i, 19, t)
for i, t in enumerate(["S", "N"]):
    put(L, 4 + i, 21, t)
# más listas en fila 12+
header(L, 12, ["TIPO_CAMBIO", "", "", "ESTADO_EMPRESARIAL", "", "", "ATRACTIVIDAD", "", "", "", "AMBITO_ALERTA", "", "DIMENSION", "", "POLARIDAD"])
for i, t in enumerate(["REAL", "METODOLÓGICO"]):
    put(L, 13 + i, 1, t)
for i, (t, _) in enumerate(M.ESCALA_ESTADO):
    put(L, 13 + i, 4, t)
for i, (t, _) in enumerate(M.ESCALA_ATRACTIVIDAD):
    put(L, 13 + i, 7, t)
AMBITOS = ["Contratos", "Cancelaciones", "Retrasos", "Sobrecostes", "Cartera", "Márgenes", "Caja", "Deuda", "Plantilla",
           "Inversiones", "Capacidad", "Exportaciones", "Programas", "Regulación y sanciones", "Presupuestos"]
for i, t in enumerate(AMBITOS):
    put(L, 13 + i, 11, t)
for i, t in enumerate(M.DIMENSIONES):
    put(L, 13 + i, 13, t)
for i, t in enumerate(["Mayor es mejor", "Menor es mejor", "Óptimo intermedio"]):
    put(L, 13 + i, 15, t)

R_CONF = "'16_LISTAS'!$A$4:$A$8"
R_NIV = "'16_LISTAS'!$D$4:$D$9"
R_NIVNUM = "'16_LISTAS'!$E$4:$E$9"
R_BMIN = "'16_LISTAS'!$G$4:$G$8"
R_BNOM = "'16_LISTAS'!$H$4:$H$8"
R_TIPO = "'16_LISTAS'!$K$4:$K$8"
R_PI = "'16_LISTAS'!$M$4:$M$6"
R_EST = "'16_LISTAS'!$O$4:$O$7"
R_CAL = "'16_LISTAS'!$Q$4:$Q$8"
R_ESC = "'16_LISTAS'!$S$4:$S$6"
R_HOR = "'16_LISTAS'!$S$7:$S$9"
R_SN = "'16_LISTAS'!$U$4:$U$5"
R_TC = "'16_LISTAS'!$A$13:$A$14"
R_ESTEMP = "'16_LISTAS'!$D$13:$D$17"
R_ATR = "'16_LISTAS'!$G$13:$G$17"
R_AMB = f"'16_LISTAS'!$K$13:$K${12 + len(AMBITOS)}"
R_DIM = "'16_LISTAS'!$M$13:$M$16"
R_IDS = f"'02_INDICADORES'!$A$2:$A${LAST}"

# =====================================================================================
# 00_README
# =====================================================================================
ws = sheet("00_README", [26, 110], freeze=None)
title(ws, f"DCIP {M.VERSION} — ARQUITECTURA MAESTRA v{M.VERSION_ARCH}", f"Versión {M.FECHA_VERSION}. Perímetro de aplicación: empresas de defensa y tecnologías estratégicas de cualquier país.")
rows = [
    ("Propósito", "Arquitectura maestra reutilizable para cualquier empresa de defensa o tecnologías estratégicas, de cualquier jurisdicción."),
    ("Separación", "Metodología, indicadores, scoring, evidencia, auditoría, ficha, histórico y estándar visual están desacoplados de cualquier empresa."),
    ("Empresa", "Se introduce solo en 01_CONFIG. Un archivo por empresa y versión de expediente; el histórico multiempresa vive en 14_HISTORICO."),
    ("Fórmula", "DCIP = 25 % FIN + 25 % IND + 25 % DEF + 25 % PRO. Pesos fijos. Cada dimensión suma 100 puntos repartidos en sus indicadores."),
    ("Indicadores", f"{N_IND} indicadores (FIN 10, IND 10, DEF 12, PRO 9) con rúbrica de cinco niveles. 100 es siempre la situación más favorable."),
    ("Rango", "Se calcula una nota normalizada (sobre pesos disponibles) y una nota prudente (indicadores sin evidencia a 25). Si cruzan una banda, la categoría se publica como rango."),
    ("Cobertura", "≥ 80 % concluyente; 60–79 % con reservas y confianza máxima MEDIA; < 60 % dimensión no concluyente y sin DCIP global."),
    ("Evidencia", "HECHO / ESTIMACIÓN / INFERENCIA / EXPECTATIVA / HIPÓTESIS. Cada puntuación cita EVIDENCIA_ID; cada evidencia cita FUENTE_ID."),
    ("Confianza", "Se deriva del nivel de fuente (N1–N6) y de la antigüedad; nunca de la puntuación. La confianza introducida no puede superar la máxima calculada."),
    ("Internacional", "01_CONFIG registra jurisdicción, norma contable, moneda, tipo de cambio y cierre fiscal. 17_JURISDICCIONES mapea las fuentes N1 y N3 por país."),
    ("Principio", "No inventar, no imputar huecos, separar perímetros y buscar contradicciones. La calidad de la conclusión nunca supera la calidad de la evidencia."),
    ("", ""),
    ("LEYENDA DE CELDAS", ""),
    ("Texto azul sobre fondo amarillo", "Celda de entrada. Es lo único que rellena el analista."),
    ("Texto negro", "Fórmula. No editar."),
    ("Texto verde", "Enlace a otra hoja. No editar."),
    ("Fondo gris y texto 'EJEMPLO'", "Fila de ejemplo con valores ficticios para mostrar el formato. Borrar antes de usar el expediente."),
    ("", ""),
    ("ORDEN DE USO", "1) 01_CONFIG → 2) 12_FUENTES → 3) 05_EVIDENCIA → 4) 07_FICHA (snapshot y ratios) → 5) 03_SCORING → 6) 04_RESULTADO → 7) 08–11 → 8) 06_AUDITORIA → 9) 13_CAMBIOS y 14_HISTORICO → 10) estándar visual (20–24)."),
    ("Hojas", "00 README · 01 CONFIG · 02 INDICADORES · 03 SCORING · 04 RESULTADO · 05 EVIDENCIA · 06 AUDITORIA · 07 FICHA · 08 RIESGOS · 09 OPORTUNIDADES · 10 PROSPECTIVA · 11 ALERTA_TEMPRANA · 12 FUENTES · 13 CAMBIOS · 14 HISTORICO · 15 RANKING · 16 LISTAS · 17 JURISDICCIONES · 20–24 estándar visual."),
]
for i, (a, b) in enumerate(rows, 4):
    put(ws, i, 1, a, font=F_BOLD, border=False)
    put(ws, i, 2, b, border=False)
put(ws, 17, 1, "Texto azul sobre fondo amarillo", font=F_INPUT, fill=FILL_INPUT, border=False)
put(ws, 19, 1, "Texto verde", font=F_LINK, border=False)
put(ws, 20, 1, "Fondo gris y texto 'EJEMPLO'", fill=FILL_EX, border=False)

# =====================================================================================
# 01_CONFIG
# =====================================================================================
C = sheet("01_CONFIG", [30, 36, 70], freeze=None)
title(C, "01_CONFIG — Parámetros del expediente", "Rellenar las celdas amarillas. Las demás hojas leen de aquí.")
header(C, 3, ["CAMPO", "VALOR", "NOTA"])
config = [
    ("EMPRESA", "", "Nombre oficial de la entidad analizada."),
    ("PAIS", "", "País de sede y control efectivo."),
    ("JURISDICCION", "", "Elegir de 17_JURISDICCIONES. Determina las fuentes N1 y el tope de confianza."),
    ("FECHA_CORTE", "", "Fecha hasta la que se considera válida la información (DD/MM/AAAA)."),
    ("PERIODO_FINANCIERO", "", "Último ejercicio cerrado y auditado a la fecha de corte (AAAA o AAAA/AA)."),
    ("CIERRE_FISCAL", "", "Mes de cierre del ejercicio (p. ej. 31/12, 31/03, 30/06)."),
    ("NORMA_CONTABLE", "", "IFRS, US GAAP, GAAP local (indicar cuál)."),
    ("MONEDA_ORIGINAL", "", "Código ISO (EUR, USD, GBP, KRW, ILS...)."),
    ("TIPO_CAMBIO_EUR", 1, "Euros por una unidad de moneda original en la fecha de corte. EUR = 1."),
    ("FUENTE_TIPO_CAMBIO", "", "BCE, tipo de referencia, fecha."),
    ("PERIMETRO_PRINCIPAL", "", "Grupo consolidado o sociedad individual."),
    ("PERIMETRO_SECUNDARIO", "", "Filiales, segmentos o programas relevantes, si procede."),
    ("PESO_DEFENSA_INGRESOS", "", "Porcentaje de ingresos de defensa sobre el total, como fracción (0,65 = 65 %)."),
    ("PESO_DEFENSA_TIPO", "", "HECHO si procede de un segmento publicado; ESTIMACIÓN si es cálculo propio."),
    ("FIN_PESO", 0.25, "Fijo por metodología."),
    ("IND_PESO", 0.25, "Fijo por metodología."),
    ("DEF_PESO", 0.25, "Fijo por metodología."),
    ("PRO_PESO", 0.25, "Fijo por metodología."),
    ("CONTROL_PESOS", "=SUM(B18:B21)", "Debe ser 1."),
    ("VERSION_METODOLOGIA", M.VERSION, "No editar."),
    ("VERSION_EXPEDIENTE", "v1", "v1, v2... Se incrementa en cada actualización y se registra en 13_CAMBIOS."),
    ("ANALISTA", "", ""),
    ("FECHA_ANALISIS", "", "Fecha de elaboración."),
]
for i, (k, v, n) in enumerate(config, 4):
    put(C, i, 1, k, font=F_BOLD)
    is_input = k not in ("FIN_PESO", "IND_PESO", "DEF_PESO", "PRO_PESO", "CONTROL_PESOS", "VERSION_METODOLOGIA")
    put(C, i, 2, v, font=F_INPUT if is_input else F_BASE, fill=FILL_INPUT if is_input else None)
    put(C, i, 3, n, font=F_NOTE)
C["B7"].number_format = "DD/MM/YYYY"
C["B26"].number_format = "DD/MM/YYYY"
C["B16"].number_format = "0.0%"
for r in (18, 19, 20, 21, 22):
    C[f"B{r}"].number_format = "0%"
dv_list(C, f"'17_JURISDICCIONES'!$A$4:$A${3 + len(M.JURISDICCIONES)}", "B6")
dv_list(C, R_TIPO, "B17")
CFG = {k: f"'01_CONFIG'!$B${i}" for i, (k, _, _) in enumerate(config, 4)}

# =====================================================================================
# 02_INDICADORES
# =====================================================================================
I = sheet("02_INDICADORES", [6, 6, 30, 6, 16, 44, 44, 20, 8, 16, 34, 34, 34, 34, 34, 50], freeze="D2")
header(I, 1, ["ID", "DIM", "INDICADOR", "PESO", "POLARIDAD", "DEFINICION", "METRICA", "HORIZONTE", "FUENTE_PREF", "NO_DOBLE_CONTEO",
              "RUBRICA_0", "RUBRICA_25", "RUBRICA_50", "RUBRICA_75", "RUBRICA_100", "NOTA"])
for r, ind in enumerate(M.INDICADORES, 2):
    f = dim_fill(ind["dim"])
    put(I, r, 1, ind["id"], font=F_BOLD, fill=f)
    put(I, r, 2, ind["dim"], font=F_BOLD, fill=f)
    put(I, r, 3, ind["nombre"], font=F_BOLD)
    put(I, r, 4, ind["peso"], align=CENTER)
    put(I, r, 5, ind["polaridad"])
    put(I, r, 6, ind["definicion"])
    put(I, r, 7, ind["metrica"])
    put(I, r, 8, ind["horizonte"])
    put(I, r, 9, ind["fuente"], align=CENTER)
    put(I, r, 10, ", ".join(ind["no_doble"]))
    for k, lvl in enumerate((0, 25, 50, 75, 100)):
        put(I, r, 11 + k, ind["rubrica"][lvl])
    put(I, r, 16, ind["nota"], font=F_NOTE)
    I.row_dimensions[r].height = 92
I.auto_filter.ref = f"A1:P{LAST}"

# =====================================================================================
# 03_SCORING
# =====================================================================================
S = sheet("03_SCORING", [6, 6, 30, 6, 11, 9, 9, 9, 9, 10, 10, 13, 10, 11, 9, 11, 11, 12, 48, 16, 30, 30, 9], freeze="D2")
cols = ["ID", "DIM", "INDICADOR", "PESO", "PUNTUACION", "RANGO_MIN", "RANGO_MAX", "DISPONIBLE", "PESO_DISP", "PONDERADA",
        "NIVEL_RUBRICA", "TIPO", "NIVEL_FUENTE", "ANTIG_MESES", "ES_CUENTAS", "CONF_MAX", "CONFIANZA", "CONTROL",
        "JUSTIFICACION", "EVIDENCIA_IDS", "CONTRADICCION", "OBSERVACIONES", "CONF_NUM"]
header(S, 1, cols)
for r in range(2, LAST + 1):
    f = dim_fill(M.INDICADORES[r - 2]["dim"])
    put(S, r, 1, f"='02_INDICADORES'!A{r}", font=F_LINK, fill=f)
    put(S, r, 2, f"='02_INDICADORES'!B{r}", font=F_LINK, fill=f)
    put(S, r, 3, f"='02_INDICADORES'!C{r}", font=F_LINK)
    put(S, r, 4, f"='02_INDICADORES'!D{r}", font=F_LINK, align=CENTER)
    for c in (5, 6, 7):
        put(S, r, c, None, font=F_INPUT, fill=FILL_INPUT, align=CENTER)
    put(S, r, 8, f'=IF(E{r}="",0,1)', align=CENTER)
    put(S, r, 9, f"=D{r}*H{r}", align=CENTER)
    put(S, r, 10, f"=IF(H{r}=1,E{r}*D{r},0)", align=CENTER)
    put(S, r, 11, f'=IF(E{r}="","",IF(E{r}<=12,0,IF(E{r}<=37,25,IF(E{r}<=62,50,IF(E{r}<=87,75,100)))))', align=CENTER)
    put(S, r, 12, None, font=F_INPUT, fill=FILL_INPUT, align=CENTER)
    put(S, r, 13, None, font=F_INPUT, fill=FILL_INPUT, align=CENTER)
    put(S, r, 14, None, font=F_INPUT, fill=FILL_INPUT, align=CENTER)
    put(S, r, 15, None, font=F_INPUT, fill=FILL_INPUT, align=CENTER)
    put(S, r, 16, (f'=IF(M{r}="","",INDEX({R_CONF},MAX(1,INDEX({R_NIVNUM},MATCH(M{r},{R_NIV},0))'
                   f'-IF(N{r}="",0,IF(N{r}<=IF(O{r}="S",18,12),0,IF(N{r}<=IF(O{r}="S",30,24),1,2))))))'), align=CENTER)
    put(S, r, 17, None, font=F_INPUT, fill=FILL_INPUT, align=CENTER)
    put(S, r, 18, (f'=IF(E{r}="","",IF(Q{r}="","FALTA CONFIANZA",IF(P{r}="","FALTA NIVEL FUENTE",'
                   f'IF(MATCH(Q{r},{R_CONF},0)>MATCH(P{r},{R_CONF},0),"EXCEDE CONF_MAX","OK"))))'), align=CENTER)
    for c in (19, 20, 21, 22):
        put(S, r, c, None, font=F_INPUT, fill=FILL_INPUT)
    put(S, r, 23, f'=IF(Q{r}="",0,MATCH(Q{r},{R_CONF},0))', align=CENTER, font=F_NOTE)
dv_num(S, f"E2:G{LAST}")
dv_list(S, R_TIPO, f"L2:L{LAST}")
dv_list(S, R_NIV, f"M2:M{LAST}")
dv_list(S, R_SN, f"O2:O{LAST}")
dv_list(S, R_CONF, f"Q2:Q{LAST}")
dv_num(S, f"N2:N{LAST}", 0, 240)
# formato condicional por bandas en PUNTUACION
for lo, hi, nom, col, fillc in M.BANDAS:
    S.conditional_formatting.add(f"E2:E{LAST}", CellIsRule(operator="between", formula=[str(lo), str(hi)], fill=PatternFill("solid", fgColor=fillc), font=Font(name=FONT, size=10, bold=True, color=col)))
S.conditional_formatting.add(f"R2:R{LAST}", FormulaRule(formula=[f'AND(R2<>"",R2<>"OK")'], fill=PatternFill("solid", fgColor="FFCDD2"), font=Font(name=FONT, size=10, bold=True, color="B71C1C")))
S["E1"].comment = Comment("Puntuación central 0–100 según la rúbrica de 02_INDICADORES. Dejar vacía si no hay evidencia suficiente: cuenta como no disponible y reduce la cobertura.", "DCIP")
S["F1"].comment = Comment("Opcional. Si la evidencia solo permite un intervalo, anotar mínimo y máximo y poner en PUNTUACION el punto central.", "DCIP")
S["N1"].comment = Comment("Meses entre la fecha de la evidencia principal y la fecha de corte. Determina la penalización de confianza.", "DCIP")
S["O1"].comment = Comment("S si la evidencia principal son cuentas anuales auditadas (umbrales de antigüedad 18/30 meses en lugar de 12/24).", "DCIP")
S["Q1"].comment = Comment("Confianza asignada. No puede superar CONF_MAX; la columna CONTROL lo señala.", "DCIP")
S.auto_filter.ref = f"A1:W{LAST}"

# =====================================================================================
# 04_RESULTADO
# =====================================================================================
R = sheet("04_RESULTADO", [16, 12, 14, 12, 18, 14, 12, 14, 12, 16, 16, 18, 30, 10], freeze=None)
title(R, "04_RESULTADO — Agregación por dimensión y DCIP global", "Todo son fórmulas. La nota se publica como rango prudente–normalizado cuando hay indicadores sin evidencia.")
header(R, 3, ["DIMENSION", "PESO_TOTAL", "PESO_DISPONIBLE", "COBERTURA", "ESTADO_COBERTURA", "NORMALIZADA", "PRUDENTE", "NOTA_DIMENSION", "NOTA_MIN",
              "CONF_MAX_COB", "CONF_INDICADORES", "CONFIANZA_DIM", "CATEGORIA", "CONF_NUM"])
SB = f"'03_SCORING'!$B$2:$B${LAST}"
SD = f"'03_SCORING'!$D$2:$D${LAST}"
SI = f"'03_SCORING'!$I$2:$I${LAST}"
SJ = f"'03_SCORING'!$J$2:$J${LAST}"
SW = f"'03_SCORING'!$W$2:$W${LAST}"
for i, d in enumerate(M.DIMENSIONES):
    r = 4 + i
    put(R, r, 1, d, font=F_BOLD, fill=dim_fill(d), align=CENTER)
    put(R, r, 2, f"=SUMIFS({SD},{SB},A{r})", align=CENTER)
    put(R, r, 3, f"=SUMIFS({SI},{SB},A{r})", align=CENTER)
    put(R, r, 4, f"=IF(B{r}=0,0,C{r}/B{r})", align=CENTER, fmt="0%")
    put(R, r, 5, f'=IF(D{r}<0.6,"NO CONCLUYENTE",IF(D{r}<0.8,"CON RESERVAS","CONCLUYENTE"))', align=CENTER)
    put(R, r, 6, f'=IF(C{r}=0,"",SUMIFS({SJ},{SB},A{r})/C{r})', align=CENTER, fmt="0.0")
    put(R, r, 7, f'=IF(B{r}=0,"",(SUMIFS({SJ},{SB},A{r})+{M.PUNT_PRUDENTE_FALTANTE}*(B{r}-C{r}))/B{r})', align=CENTER, fmt="0.0")
    put(R, r, 8, f'=IF(E{r}="NO CONCLUYENTE","NC",ROUND(F{r},0))', align=CENTER, font=F_BOLD)
    put(R, r, 9, f'=IF(E{r}="NO CONCLUYENTE","NC",ROUND(G{r},0))', align=CENTER)
    put(R, r, 10, f'=IF(D{r}<0.6,"",IF(D{r}<0.8,"MEDIA","MUY ALTA"))', align=CENTER)
    put(R, r, 11, f'=IF(C{r}=0,"",SUMPRODUCT(({SB}=A{r})*({SI})*({SW}))/C{r})', align=CENTER, fmt="0.00")
    put(R, r, 12, (f'=IF(E{r}="NO CONCLUYENTE","NO CONCLUYENTE",IF(K{r}="","",INDEX({R_CONF},'
                   f'MAX(1,MIN(IF(D{r}<0.8,3,5),ROUNDDOWN(K{r},0))))))'), align=CENTER, font=F_BOLD)
    put(R, r, 13, (f'=IF(H{r}="NC","NO CONCLUYENTE",IF(LOOKUP(H{r},{R_BMIN},{R_BNOM})=LOOKUP(I{r},{R_BMIN},{R_BNOM}),'
                   f'LOOKUP(H{r},{R_BMIN},{R_BNOM}),LOOKUP(I{r},{R_BMIN},{R_BNOM})&" – "&LOOKUP(H{r},{R_BMIN},{R_BNOM})&" (rango)"))'), align=CENTER)
    put(R, r, 14, f'=IF(L{r}="NO CONCLUYENTE",0,IF(L{r}="",0,MATCH(L{r},{R_CONF},0)))', align=CENTER, font=F_NOTE)
# DCIP global
put(R, 9, 1, "DCIP GLOBAL", font=F_BOLD, fill=FILL_GREY)
header(R, 10, ["CONCEPTO", "VALOR", "NOTA"])
pesos = f"{CFG['FIN_PESO']},{CFG['IND_PESO']},{CFG['DEF_PESO']},{CFG['PRO_PESO']}"
glob = [
    ("DCIP_NORMALIZADO", f'=IF(COUNTIF(E4:E7,"NO CONCLUYENTE")>0,"NO CONCLUYENTE",ROUND(F4*{CFG["FIN_PESO"]}+F5*{CFG["IND_PESO"]}+F6*{CFG["DEF_PESO"]}+F7*{CFG["PRO_PESO"]},0))', "Media ponderada de las notas normalizadas, redondeada al entero."),
    ("DCIP_PRUDENTE", f'=IF(COUNTIF(E4:E7,"NO CONCLUYENTE")>0,"NO CONCLUYENTE",ROUND(G4*{CFG["FIN_PESO"]}+G5*{CFG["IND_PESO"]}+G6*{CFG["DEF_PESO"]}+G7*{CFG["PRO_PESO"]},0))', "Límite inferior del rango: indicadores sin evidencia a 25."),
    ("CATEGORIA_DCIP", (f'=IF(B11="NO CONCLUYENTE","NO CONCLUYENTE",IF(LOOKUP(B11,{R_BMIN},{R_BNOM})=LOOKUP(B12,{R_BMIN},{R_BNOM}),'
                        f'LOOKUP(B11,{R_BMIN},{R_BNOM}),LOOKUP(B12,{R_BMIN},{R_BNOM})&" – "&LOOKUP(B11,{R_BMIN},{R_BNOM})&" (rango)"))'), "Si el rango cruza una banda, se publica el rango."),
    ("CONFIANZA_GLOBAL", f'=IF(MIN(N4:N7)=0,"NO CONCLUYENTE",INDEX({R_CONF},MIN(N4:N7)))', "La menor de las cuatro dimensiones."),
    ("COBERTURA_GLOBAL", "=AVERAGE(D4:D7)", "Media simple de coberturas."),
    ("CONTROL_PESOS", f'=IF(ROUND({CFG["CONTROL_PESOS"]},6)=1,"OK","ERROR: los pesos no suman 1")', ""),
    ("CONTROLES_SCORING", f'=COUNTIF(\'03_SCORING\'!R2:R{LAST},"OK")&" OK / "&(COUNTA(\'03_SCORING\'!E2:E{LAST})-COUNTIF(\'03_SCORING\'!R2:R{LAST},"OK"))&" con incidencia"', "Ver columna CONTROL de 03_SCORING."),
    ("ESTADO_EMPRESARIAL", "", "Juicio del analista, escala de 16_LISTAS. Informado por FIN e IND; no se calcula."),
    ("ATRACTIVIDAD_ESTRATEGICA", "", "Juicio del analista, escala de 16_LISTAS. Informado por DEF y PRO; no se calcula."),
    ("VALIDEZ_AUDITORIA", "='06_AUDITORIA'!B4", "Si es NO VÁLIDA, el expediente no se publica."),
]
for i, (k, v, n) in enumerate(glob, 11):
    put(R, i, 1, k, font=F_BOLD)
    is_in = k in ("ESTADO_EMPRESARIAL", "ATRACTIVIDAD_ESTRATEGICA")
    put(R, i, 2, v, font=F_INPUT if is_in else F_BOLD, fill=FILL_INPUT if is_in else None, align=CENTER)
    put(R, i, 3, n, font=F_NOTE)
R["B15"].number_format = "0%"
dv_list(R, R_ESTEMP, "B18")
dv_list(R, R_ATR, "B19")
for lo, hi, nom, col, fillc in M.BANDAS:
    for rng in ("H4:I7", "B11:B12"):
        R.conditional_formatting.add(rng, CellIsRule(operator="between", formula=[str(lo), str(hi)], fill=PatternFill("solid", fgColor=fillc), font=Font(name=FONT, size=10, bold=True, color=col)))
R.conditional_formatting.add("E4:E7", FormulaRule(formula=['E4="NO CONCLUYENTE"'], fill=PatternFill("solid", fgColor="FFCDD2")))
R.conditional_formatting.add("E4:E7", FormulaRule(formula=['E4="CON RESERVAS"'], fill=PatternFill("solid", fgColor="FFE0B2")))
# tabla de bandas y escalas de juicio para referencia
put(R, 23, 1, "BANDAS DCIP", font=F_BOLD, fill=FILL_GREY)
header(R, 24, ["BANDA", "MIN", "MAX"])
for i, (lo, hi, nom, col, fillc) in enumerate(M.BANDAS, 25):
    put(R, i, 1, nom, fill=PatternFill("solid", fgColor=fillc), font=Font(name=FONT, size=10, bold=True, color=col))
    put(R, i, 2, lo, align=CENTER)
    put(R, i, 3, hi, align=CENTER)
put(R, 31, 1, "ESCALAS DE JUICIO", font=F_BOLD, fill=FILL_GREY)
header(R, 32, ["ESTADO EMPRESARIAL", "", "DESCRIPCION"])
for i, (t, d) in enumerate(M.ESCALA_ESTADO, 33):
    put(R, i, 1, t, font=F_BOLD)
    R.merge_cells(start_row=i, start_column=3, end_row=i, end_column=13)
    put(R, i, 3, d)
header(R, 39, ["ATRACTIVIDAD ESTRATEGICA", "", "DESCRIPCION"])
for i, (t, d) in enumerate(M.ESCALA_ATRACTIVIDAD, 40):
    put(R, i, 1, t, font=F_BOLD)
    R.merge_cells(start_row=i, start_column=3, end_row=i, end_column=13)
    put(R, i, 3, d)

# =====================================================================================
# 05_EVIDENCIA
# =====================================================================================
E = sheet("05_EVIDENCIA", [12, 10, 18, 12, 12, 10, 12, 10, 9, 9, 12, 50, 24, 16, 9, 9, 10, 11, 24, 24, 24], freeze="C2")
ecols = ["EVIDENCIA_ID", "ID_INDICADOR", "EMPRESA", "FECHA_CORTE", "FECHA_EVIDENCIA", "ANTIG_MESES", "POSTERIOR_CORTE", "FUENTE_ID", "NIVEL_FUENTE",
         "ES_CUENTAS", "TIPO", "CITA_EXTRACTO", "REFERENCIA (página / URL)", "PERIMETRO", "MONEDA", "IDIOMA", "TRADUCCION", "CONF_MAX",
         "CONTRADICCION", "RESOLUCION", "OBSERVACIONES"]
header(E, 1, ecols)
NE = 300
FID = "'12_FUENTES'!$A$2:$A$301"
FNIV = "'12_FUENTES'!$B$2:$B$301"
for r in range(2, NE + 2):
    for c in (1, 2, 5, 8, 10, 11, 12, 13, 14, 15, 16, 17, 19, 20, 21):
        put(E, r, c, None, font=F_INPUT, fill=FILL_INPUT)
    put(E, r, 3, f'=IF(A{r}="","",{CFG["EMPRESA"]})', font=F_LINK)
    put(E, r, 4, f'=IF(A{r}="","",{CFG["FECHA_CORTE"]})', font=F_LINK, fmt="DD/MM/YYYY")
    E[f"E{r}"].number_format = "DD/MM/YYYY"
    put(E, r, 6, f'=IF(OR(E{r}="",D{r}=""),"",(YEAR(D{r})-YEAR(E{r}))*12+MONTH(D{r})-MONTH(E{r}))', align=CENTER)
    put(E, r, 7, f'=IF(OR(E{r}="",D{r}=""),"",IF(E{r}>D{r},"EXCLUIR","OK"))', align=CENTER)
    put(E, r, 9, f'=IF(H{r}="","",IFERROR(INDEX({FNIV},MATCH(H{r},{FID},0)),"FUENTE NO REGISTRADA"))', font=F_LINK, align=CENTER)
    put(E, r, 18, (f'=IF(OR(I{r}="",I{r}="FUENTE NO REGISTRADA"),"",INDEX({R_CONF},MAX(1,INDEX({R_NIVNUM},MATCH(I{r},{R_NIV},0))'
                   f'-IF(F{r}="",0,IF(F{r}<=IF(J{r}="S",18,12),0,IF(F{r}<=IF(J{r}="S",30,24),1,2))))))'), align=CENTER)
dv_list(E, R_IDS, f"B2:B{NE + 1}")
dv_list(E, R_SN, f"J2:J{NE + 1}")
dv_list(E, R_TIPO, f"K2:K{NE + 1}")
dv_list(E, R_SN, f"Q2:Q{NE + 1}")
E.conditional_formatting.add(f"G2:G{NE + 1}", FormulaRule(formula=['G2="EXCLUIR"'], fill=PatternFill("solid", fgColor="FFCDD2"), font=Font(name=FONT, size=10, bold=True, color="B71C1C")))
E.conditional_formatting.add(f"I2:I{NE + 1}", FormulaRule(formula=['I2="FUENTE NO REGISTRADA"'], fill=PatternFill("solid", fgColor="FFCDD2")))
# fila de ejemplo
ex = ["E001", "F01", None, None, dt.date(2026, 3, 31), None, None, "F001", None, "S", "HECHO",
      "EJEMPLO — Resultado de explotación 2025: 184,2 M (moneda local). Cifra de negocio: 2.310,5 M.",
      "Cuentas anuales consolidadas 2025, p. 47, cuenta de resultados", "Grupo consolidado", "EUR", "es", "N", None, "", "", "EJEMPLO. Borrar antes de usar."]
for c, v in enumerate(ex, 1):
    if v is not None:
        E.cell(row=2, column=c, value=v)
    E.cell(row=2, column=c).fill = FILL_EX
    E.cell(row=2, column=c).font = Font(name=FONT, size=10, italic=True, color="616161")
E.auto_filter.ref = f"A1:U{NE + 1}"

# =====================================================================================
# 06_AUDITORIA
# =====================================================================================
A = sheet("06_AUDITORIA", [10, 46, 9, 60, 14, 30, 30], freeze="A7")
title(A, "06_AUDITORIA — Controles de calidad del expediente", "Si un control CRÍTICO está en NO CUMPLE, el expediente es NO VÁLIDO y no se publica.")
put(A, 3, 1, "CRÍTICOS NO CUMPLE", font=F_BOLD); put(A, 3, 2, '=COUNTIFS(C7:C40,"S",E7:E40,"NO CUMPLE")', align=CENTER, font=F_BOLD)
put(A, 4, 1, "VALIDEZ", font=F_BOLD); put(A, 4, 2, '=IF(B3>0,"NO VÁLIDA",IF(COUNTIF(E7:E40,"PENDIENTE")>0,"EN CURSO","VÁLIDA"))', align=CENTER, font=F_BOLD)
put(A, 5, 1, "PENDIENTES", font=F_BOLD); put(A, 5, 2, '=COUNTIF(E7:E40,"PENDIENTE")', align=CENTER)
header(A, 6, ["CONTROL_ID", "CONTROL", "CRITICO", "METODO_COMPROBACION", "ESTADO", "EVIDENCIA", "OBSERVACIONES"])
for i, (cid, ctrl, crit, met) in enumerate(M.CONTROLES_AUDITORIA, 7):
    put(A, i, 1, cid, align=CENTER)
    put(A, i, 2, ctrl, font=F_BOLD if crit == "S" else F_BASE)
    put(A, i, 3, crit, align=CENTER)
    put(A, i, 4, met, font=F_NOTE)
    put(A, i, 5, "PENDIENTE", font=F_INPUT, fill=FILL_INPUT, align=CENTER)
    put(A, i, 6, None, font=F_INPUT, fill=FILL_INPUT)
    put(A, i, 7, None, font=F_INPUT, fill=FILL_INPUT)
dv_list(A, R_EST, "E7:E40")
A.conditional_formatting.add("B4", FormulaRule(formula=['B4="NO VÁLIDA"'], fill=PatternFill("solid", fgColor="FFCDD2")))
A.conditional_formatting.add("B4", FormulaRule(formula=['B4="VÁLIDA"'], fill=PatternFill("solid", fgColor="C8E6C9")))
A.conditional_formatting.add("E7:E40", FormulaRule(formula=['E7="NO CUMPLE"'], fill=PatternFill("solid", fgColor="FFCDD2")))
A.conditional_formatting.add("E7:E40", FormulaRule(formula=['E7="CUMPLE"'], fill=PatternFill("solid", fgColor="C8E6C9")))

# =====================================================================================
# 07_FICHA (snapshot + ratios + texto)
# =====================================================================================
F = sheet("07_FICHA", [34, 18, 16, 14, 14, 60], freeze=None)
title(F, "07_FICHA — Snapshot, ratios y ficha empresarial", "Importes en moneda original; conversión a M€ automática con TIPO_CAMBIO_EUR de 01_CONFIG. Los ratios alimentan las rúbricas FIN.")
header(F, 3, ["MAGNITUD", "VALOR (moneda original, M)", "VALOR (M€)", "EVIDENCIA_ID", "TIPO", "NOTA"])
snap = [
    ("Cifra de negocio", "Ingresos reconocidos del PERIODO_FINANCIERO."),
    ("Ingresos de defensa", "Segmento publicado o estimación; etiquetar."),
    ("Contratación (order intake)", "No equivale a ingresos."),
    ("Cartera total (backlog)", "Según definición de la fuente."),
    ("Cartera firme", "Solo compromisos contractuales firmes."),
    ("Resultado de explotación (EBIT)", "Reportado, no ajustado."),
    ("EBITDA", "Solo publicado o cálculo propio declarado."),
    ("Resultado neto", ""),
    ("Flujo de caja operativo", ""),
    ("Caja y equivalentes", ""),
    ("Deuda financiera bruta", "Préstamos + bonos + arrendamientos."),
    ("Instrumentos híbridos / participativos", "Separados de la deuda ordinaria."),
    ("Patrimonio neto", ""),
    ("Activo total", ""),
    ("Activo corriente", ""),
    ("Pasivo corriente", ""),
    ("Capex", "Inversión en inmovilizado material e intangible."),
    ("Amortización", ""),
    ("Apoyo público del ejercicio", "Según definición de F07."),
    ("Provisiones y contingencias extraordinarias", "Según definición de F10."),
    ("Plantilla (FTE)", "Número, no importe."),
    ("Ingresos de exportación de defensa", ""),
]
SNAP_ROW = {}
for i, (k, n) in enumerate(snap, 4):
    SNAP_ROW[k] = i
    put(F, i, 1, k, font=F_BOLD)
    put(F, i, 2, None, font=F_INPUT, fill=FILL_INPUT, fmt="#,##0.0")
    if k == "Plantilla (FTE)":
        put(F, i, 3, f"=B{i}", fmt="#,##0")
    else:
        put(F, i, 3, f'=IF(B{i}="","",B{i}*{CFG["TIPO_CAMBIO_EUR"]})', fmt="#,##0.0")
    put(F, i, 4, None, font=F_INPUT, fill=FILL_INPUT)
    put(F, i, 5, None, font=F_INPUT, fill=FILL_INPUT)
    put(F, i, 6, n, font=F_NOTE)
dv_list(F, R_TIPO, f"E4:E{3 + len(snap)}")
rr = 4 + len(snap) + 1
put(F, rr, 1, "RATIOS CALCULADOS (moneda original; alimentan las rúbricas FIN)", font=F_BOLD, fill=FILL_GREY)
F.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=6)
header(F, rr + 1, ["RATIO", "VALOR", "INDICADOR", "", "", "DEFINICION"])


def sr(k):
    return f"B{SNAP_ROW[k]}"


ratios = [
    ("Margen EBIT", f'=IF(OR({sr("Cifra de negocio")}="",{sr("Resultado de explotación (EBIT)")}=""),"",{sr("Resultado de explotación (EBIT)")}/{sr("Cifra de negocio")})', "F01", "0.0%", "EBIT / cifra de negocio."),
    ("FCO / EBITDA", f'=IF(OR({sr("EBITDA")}="",{sr("Flujo de caja operativo")}="",{sr("EBITDA")}=0),"",{sr("Flujo de caja operativo")}/{sr("EBITDA")})', "F03", "0.0%", "Un solo ejercicio; la rúbrica usa la media de dos."),
    ("Ratio corriente", f'=IF(OR({sr("Activo corriente")}="",{sr("Pasivo corriente")}="",{sr("Pasivo corriente")}=0),"",{sr("Activo corriente")}/{sr("Pasivo corriente")})', "F04", "0.00", "Activo corriente / pasivo corriente."),
    ("Deuda financiera neta", f'=IF(OR({sr("Deuda financiera bruta")}="",{sr("Caja y equivalentes")}=""),"",{sr("Deuda financiera bruta")}-{sr("Caja y equivalentes")})', "F05", "#,##0.0", "Deuda bruta − caja. Híbridos excluidos."),
    ("Deuda neta / EBITDA", f'=IF(OR(B{rr+5}="",{sr("EBITDA")}="",{sr("EBITDA")}<=0),"",B{rr+5}/{sr("EBITDA")})', "F05", "0.00x", "Si EBITDA ≤ 0 con deuda, puntuar 0 en F05."),
    ("Patrimonio / activo", f'=IF(OR({sr("Patrimonio neto")}="",{sr("Activo total")}="",{sr("Activo total")}=0),"",{sr("Patrimonio neto")}/{sr("Activo total")})', "F06", "0.0%", ""),
    ("Apoyo público / ingresos", f'=IF(OR({sr("Apoyo público del ejercicio")}="",{sr("Cifra de negocio")}=""),"",{sr("Apoyo público del ejercicio")}/{sr("Cifra de negocio")})', "F07", "0.0%", "Un ejercicio; la rúbrica usa la media de tres."),
    ("Cartera firme / ingresos (años)", f'=IF(OR({sr("Cartera firme")}="",{sr("Cifra de negocio")}=""),"",{sr("Cartera firme")}/{sr("Cifra de negocio")})', "F08", "0.00", ""),
    ("Book-to-bill", f'=IF(OR({sr("Contratación (order intake)")}="",{sr("Cifra de negocio")}=""),"",{sr("Contratación (order intake)")}/{sr("Cifra de negocio")})', "F08", "0.00", "Contratación / ingresos."),
    ("Capex / amortización", f'=IF(OR({sr("Capex")}="",{sr("Amortización")}="",{sr("Amortización")}=0),"",{sr("Capex")}/{sr("Amortización")})', "I08", "0.00x", ""),
    ("Contingencias / ingresos", f'=IF(OR({sr("Provisiones y contingencias extraordinarias")}="",{sr("Cifra de negocio")}=""),"",{sr("Provisiones y contingencias extraordinarias")}/{sr("Cifra de negocio")})', "F10", "0.0%", ""),
    ("Peso de defensa en ingresos", f'=IF(OR({sr("Ingresos de defensa")}="",{sr("Cifra de negocio")}=""),"",{sr("Ingresos de defensa")}/{sr("Cifra de negocio")})', "CONFIG", "0.0%", "Debe coincidir con PESO_DEFENSA_INGRESOS de 01_CONFIG."),
    ("Exportación / ingresos de defensa", f'=IF(OR({sr("Ingresos de exportación de defensa")}="",{sr("Ingresos de defensa")}="",{sr("Ingresos de defensa")}=0),"",{sr("Ingresos de exportación de defensa")}/{sr("Ingresos de defensa")})', "D07", "0.0%", ""),
    ("Ingresos por empleado (miles, moneda original)", f'=IF(OR({sr("Cifra de negocio")}="",{sr("Plantilla (FTE)")}="",{sr("Plantilla (FTE)")}=0),"",{sr("Cifra de negocio")}*1000/{sr("Plantilla (FTE)")})', "I04", "#,##0", "Comparar con pares homogéneos."),
]
for i, (k, f, ind, fmt, n) in enumerate(ratios, rr + 2):
    put(F, i, 1, k, font=F_BOLD)
    put(F, i, 2, f, fmt=fmt, align=CENTER)
    put(F, i, 3, ind, align=CENTER)
    put(F, i, 6, n, font=F_NOTE)
tr = rr + 2 + len(ratios) + 1
put(F, tr, 1, "FICHA EMPRESARIAL (texto)", font=F_BOLD, fill=FILL_GREY)
F.merge_cells(start_row=tr, start_column=1, end_row=tr, end_column=6)
header(F, tr + 1, ["CAMPO", "CONTENIDO", "", "", "", "TIPO_EVIDENCIA_DOMINANTE"])
ficha = ["IDENTIFICACION", "SITUACION", "PROPIEDAD_CONTROL", "FINANZAS", "DEFENSA_PROGRAMAS", "INDUSTRIA_TECNOLOGIA", "MERCADO_INTERNACIONALIZACION",
         "COMPETENCIA", "RIESGOS", "OPORTUNIDADES", "PROSPECTIVA", "ALERTA_TEMPRANA", "DIAGNOSTICO_FIN", "DIAGNOSTICO_IND", "DIAGNOSTICO_DEF", "DIAGNOSTICO_PRO",
         "FORTALEZAS (3–5)", "VULNERABILIDADES (3–5)", "ESTADO_EMPRESARIAL", "ATRACTIVIDAD_ESTRATEGICA", "JUICIO_FINAL"]
for i, k in enumerate(ficha, tr + 2):
    put(F, i, 1, k, font=F_BOLD)
    F.merge_cells(start_row=i, start_column=2, end_row=i, end_column=5)
    put(F, i, 2, None, font=F_INPUT, fill=FILL_INPUT)
    put(F, i, 6, None, font=F_INPUT, fill=FILL_INPUT)
    F.row_dimensions[i].height = 48
dv_list(F, R_TIPO, f"F{tr + 2}:F{tr + 1 + len(ficha)}")

# =====================================================================================
# 08_RIESGOS / 09_OPORTUNIDADES
# =====================================================================================
def riesgos(name, titulo, extra):
    W = sheet(name, [8, 44, 13, 11, 12, 30, 30, 14, 14, 13, 30], freeze="A2")
    cols = ["ID", titulo, "PROBABILIDAD", "IMPACTO", "SEVERIDAD", extra, "MITIGACION / CONDICION", "HORIZONTE", "EVIDENCIA_ID", "TIPO", "OBSERVACIONES"]
    header(W, 1, cols)
    for r in range(2, 42):
        for c in (1, 2, 3, 4, 6, 7, 8, 9, 10, 11):
            put(W, r, c, None, font=F_INPUT, fill=FILL_INPUT)
        put(W, r, 5, (f'=IF(OR(C{r}="",D{r}=""),"",INDEX({{"BAJA","MEDIA","ALTA","CRÍTICA"}},'
                      f'IF(MATCH(C{r},{R_PI},0)*MATCH(D{r},{R_PI},0)>=6,4,IF(MATCH(C{r},{R_PI},0)*MATCH(D{r},{R_PI},0)>=3,3,'
                      f'IF(MATCH(C{r},{R_PI},0)*MATCH(D{r},{R_PI},0)>=2,2,1)))))'), align=CENTER, font=F_BOLD)
    dv_list(W, R_PI, "C2:D41")
    dv_list(W, R_HOR, "H2:H41")
    dv_list(W, R_TIPO, "J2:J41")
    W.conditional_formatting.add("E2:E41", FormulaRule(formula=['E2="CRÍTICA"'], fill=PatternFill("solid", fgColor="FFCDD2")))
    W.conditional_formatting.add("E2:E41", FormulaRule(formula=['E2="ALTA"'], fill=PatternFill("solid", fgColor="FFE0B2")))
    W["E1"].comment = Comment("Probabilidad × impacto (Baja=1, Media=2, Alta=3). Producto ≥6 CRÍTICA; ≥3 ALTA; 2 MEDIA; 1 BAJA. Sin probabilidades numéricas, conforme a la metodología.", "DCIP")
    return W


riesgos("08_RIESGOS", "RIESGO", "INDICADOR_ADELANTADO")
riesgos("09_OPORTUNIDADES", "OPORTUNIDAD", "CONDICION_HABILITANTE")

# =====================================================================================
# 10_PROSPECTIVA
# =====================================================================================
P = sheet("10_PROSPECTIVA", [12, 14, 44, 30, 30, 30, 30, 14, 13], freeze="A2")
header(P, 1, ["ESCENARIO", "HORIZONTE", "DESCRIPCION", "IMPULSORES", "RIESGOS", "SEÑALES_TEMPRANAS", "FACTORES_INVALIDACION", "EVIDENCIA_ID", "TIPO"])
r = 2
for h in ("12–24 meses", "3–5 años", "5–10 años"):
    for e in ("BASE", "ALCISTA", "BAJISTA"):
        put(P, r, 1, e, font=F_BOLD, align=CENTER)
        put(P, r, 2, h, align=CENTER)
        for c in range(3, 10):
            put(P, r, c, None, font=F_INPUT, fill=FILL_INPUT)
        P.row_dimensions[r].height = 60
        r += 1
dv_list(P, R_TIPO, "I2:I10")

# =====================================================================================
# 11_ALERTA_TEMPRANA
# =====================================================================================
AT = sheet("11_ALERTA_TEMPRANA", [11, 22, 34, 40, 24, 14, 34, 12], freeze="A2")
header(AT, 1, ["PRIORIDAD", "AMBITO", "INDICADOR / SEÑAL", "TRIGGER (umbral que dispara la alerta)", "FUENTE DE SEGUIMIENTO", "FRECUENCIA", "ACCION", "ESTADO"])
for r in range(2, 42):
    for c in range(1, 9):
        put(AT, r, c, None, font=F_INPUT, fill=FILL_INPUT)
dv_list(AT, R_PI, "A2:A41")
dv_list(AT, R_AMB, "B2:B41")
dv_list(AT, '"Mensual,Trimestral,Semestral,Anual,Evento"', "F2:F41")
dv_list(AT, '"ACTIVA,DISPARADA,CERRADA"', "H2:H41")

# =====================================================================================
# 12_FUENTES
# =====================================================================================
FU = sheet("12_FUENTES", [10, 8, 16, 26, 40, 12, 36, 8, 18, 14, 8, 40, 24], freeze="A2")
header(FU, 1, ["FUENTE_ID", "NIVEL", "JURISDICCION", "ORGANISMO / EMPRESA", "TITULO", "FECHA", "REFERENCIA (URL / registro)", "IDIOMA", "PERIMETRO", "NORMA_CONTABLE", "MONEDA", "AFIRMACIONES_SUSTENTADAS", "OBSERVACIONES"])
for r in range(2, 302):
    for c in range(1, 14):
        put(FU, r, c, None, font=F_INPUT, fill=FILL_INPUT)
    FU[f"F{r}"].number_format = "DD/MM/YYYY"
dv_list(FU, R_NIV, "B2:B301")
exf = ["F001", "N1", "España", "[Empresa ejemplo] / Registro Mercantil", "Cuentas anuales consolidadas e informe de auditoría 2025", dt.date(2026, 3, 31),
       "Registro Mercantil, depósito de cuentas 2025", "es", "Grupo consolidado", "IFRS", "EUR", "Snapshot financiero; F01–F10", "EJEMPLO. Borrar antes de usar."]
for c, v in enumerate(exf, 1):
    FU.cell(row=2, column=c, value=v)
    FU.cell(row=2, column=c).fill = FILL_EX
    FU.cell(row=2, column=c).font = Font(name=FONT, size=10, italic=True, color="616161")
FU.auto_filter.ref = "A1:M301"

# =====================================================================================
# 13_CAMBIOS
# =====================================================================================
CB = sheet("13_CAMBIOS", [12, 12, 14, 16, 14, 14, 10, 16, 20, 13, 40], freeze="A2")
header(CB, 1, ["VERSION_EXPEDIENTE", "FECHA", "FECHA_CORTE_NUEVA", "ELEMENTO (ID / DIM / DCIP)", "VALOR_ANTERIOR", "VALOR_NUEVO", "DELTA", "TIPO_CAMBIO", "CALIFICACION", "EVIDENCIA_ID", "EXPLICACION"])
for r in range(2, 202):
    for c in (1, 2, 3, 4, 5, 6, 8, 9, 10, 11):
        put(CB, r, c, None, font=F_INPUT, fill=FILL_INPUT)
    put(CB, r, 7, f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(F{r})),F{r}-E{r},"")', align=CENTER)
    CB[f"B{r}"].number_format = "DD/MM/YYYY"
    CB[f"C{r}"].number_format = "DD/MM/YYYY"
dv_list(CB, R_TC, "H2:H201")
dv_list(CB, R_CAL, "I2:I201")
exc = ["v2", dt.date(2026, 10, 1), dt.date(2026, 9, 30), "F05", 50, 75, None, "REAL", "MEJORA", "E014", "EJEMPLO. Deuda neta/EBITDA pasa de 2,4x a 1,6x tras amortización anticipada. Borrar antes de usar."]
for c, v in enumerate(exc, 1):
    if v is not None:
        CB.cell(row=2, column=c, value=v)
    CB.cell(row=2, column=c).fill = FILL_EX
    CB.cell(row=2, column=c).font = Font(name=FONT, size=10, italic=True, color="616161")
CB["H1"].comment = Comment("REAL: cambia la situación de la empresa. METODOLÓGICO: cambia la regla, la rúbrica o el perímetro. Nunca mezclar.", "DCIP")

# =====================================================================================
# 14_HISTORICO (formato largo, multiempresa) + 15_RANKING
# =====================================================================================
H = sheet("14_HISTORICO", [20, 14, 12, 10, 8, 6, 26, 6, 10, 10, 10, 12, 30], freeze="A2")
header(H, 1, ["EMPRESA", "PAIS", "FECHA_CORTE", "VERSION", "ID_INDICADOR", "DIM", "INDICADOR", "PESO", "PUNTUACION", "PONDERADA", "CONFIANZA", "TIPO", "OBSERVACIONES"])
NH = 2000
IDS = f"'02_INDICADORES'!$A$2:$A${LAST}"
for r in range(2, NH + 2):
    for c in (1, 2, 3, 4, 5, 9, 11, 12, 13):
        put(H, r, c, None, font=F_INPUT, fill=FILL_INPUT)
    put(H, r, 6, f'=IF(E{r}="","",IFERROR(INDEX(\'02_INDICADORES\'!$B$2:$B${LAST},MATCH(E{r},{IDS},0)),"ID?"))', font=F_LINK, align=CENTER)
    put(H, r, 7, f'=IF(E{r}="","",IFERROR(INDEX(\'02_INDICADORES\'!$C$2:$C${LAST},MATCH(E{r},{IDS},0)),"ID?"))', font=F_LINK)
    put(H, r, 8, f'=IF(E{r}="","",IFERROR(INDEX(\'02_INDICADORES\'!$D$2:$D${LAST},MATCH(E{r},{IDS},0)),""))', font=F_LINK, align=CENTER)
    put(H, r, 10, f'=IF(OR(I{r}="",H{r}=""),"",I{r}*H{r})', align=CENTER)
    H[f"C{r}"].number_format = "DD/MM/YYYY"
dv_list(H, IDS, f"E2:E{NH + 1}")
dv_list(H, R_CONF, f"K2:K{NH + 1}")
dv_list(H, R_TIPO, f"L2:L{NH + 1}")
dv_num(H, f"I2:I{NH + 1}")
# Ejemplo: dos empresas ficticias con puntuaciones por indicador (determinista, no aleatorio)
import itertools
ex_rows = []
for ei, (emp, pais, base) in enumerate((("EJEMPLO_A", "País A", 70), ("EJEMPLO_B", "País B", 45))):
    for k, ind in enumerate(M.INDICADORES):
        score = base + ((k * 7 + ei * 3) % 5) * 5 - 10  # valores plausibles entre base-10 y base+10
        score = max(0, min(100, score))
        if emp == "EJEMPLO_B" and ind["id"] in ("I07", "D08", "P04"):
            score = None  # muestra el efecto de la cobertura
        ex_rows.append((emp, pais, dt.date(2026, 6, 30), "v1", ind["id"], score, "ALTA" if ei == 0 else "MEDIA", "HECHO"))
for i, (emp, pais, fc, ver, iid, sc, conf, tipo) in enumerate(ex_rows, 2):
    vals = {1: emp, 2: pais, 3: fc, 4: ver, 5: iid, 9: sc, 11: conf, 12: tipo, 13: "EJEMPLO. Borrar antes de usar."}
    for c, v in vals.items():
        if v is not None:
            H.cell(row=i, column=c, value=v)
    for c in (1, 2, 3, 4, 5, 9, 11, 12, 13):
        H.cell(row=i, column=c).fill = FILL_EX
        H.cell(row=i, column=c).font = Font(name=FONT, size=10, italic=True, color="616161")
H.auto_filter.ref = f"A1:M{NH + 1}"

RK = sheet("15_RANKING", [22, 14, 10, 10, 10, 10, 11, 10, 10, 10, 10, 12, 20, 8], freeze="A4")
title(RK, "15_RANKING — Comparación multiempresa a partir de 14_HISTORICO", "Escribir EMPRESA y FECHA_CORTE (mismas que en 14_HISTORICO). El resto son fórmulas. Comparar solo con la misma fecha de corte y metodología.")
header(RK, 3, ["EMPRESA", "FECHA_CORTE", "COB_FIN", "COB_IND", "COB_DEF", "COB_PRO", "FIN", "IND", "DEF", "PRO", "DCIP", "CATEGORIA", "ESTADO", "RANK"])
HA = f"'14_HISTORICO'!$A$2:$A${NH + 1}"
HC = f"'14_HISTORICO'!$C$2:$C${NH + 1}"
HF = f"'14_HISTORICO'!$F$2:$F${NH + 1}"
HH = f"'14_HISTORICO'!$H$2:$H${NH + 1}"
HI = f"'14_HISTORICO'!$I$2:$I${NH + 1}"
HJ = f"'14_HISTORICO'!$J$2:$J${NH + 1}"
for r in range(4, 44):
    put(RK, r, 1, None, font=F_INPUT, fill=FILL_INPUT)
    put(RK, r, 2, None, font=F_INPUT, fill=FILL_INPUT, fmt="DD/MM/YYYY")
    for k, d in enumerate(M.DIMENSIONES):
        cob = f'=IF(A{r}="","",SUMIFS({HH},{HA},A{r},{HC},B{r},{HF},"{d}",{HI},"<>")/100)'
        put(RK, r, 3 + k, cob, align=CENTER, fmt="0%")
        colL = get_column_letter(3 + k)
        nota = (f'=IF(A{r}="","",IF({colL}{r}<0.6,"NC",ROUND(SUMIFS({HJ},{HA},A{r},{HC},B{r},{HF},"{d}")'
                f'/SUMIFS({HH},{HA},A{r},{HC},B{r},{HF},"{d}",{HI},"<>"),0)))')
        put(RK, r, 7 + k, nota, align=CENTER, font=F_BOLD)
    put(RK, r, 11, (f'=IF(A{r}="","",IF(COUNTIF(G{r}:J{r},"NC")>0,"NC",ROUND(G{r}*{CFG["FIN_PESO"]}+H{r}*{CFG["IND_PESO"]}'
                    f'+I{r}*{CFG["DEF_PESO"]}+J{r}*{CFG["PRO_PESO"]},0)))'), align=CENTER, font=F_BOLD)
    put(RK, r, 12, f'=IF(OR(K{r}="",K{r}="NC"),"",LOOKUP(K{r},{R_BMIN},{R_BNOM}))', align=CENTER)
    put(RK, r, 13, f'=IF(A{r}="","",IF(MIN(C{r}:F{r})<0.6,"NO CONCLUYENTE",IF(MIN(C{r}:F{r})<0.8,"CON RESERVAS","CONCLUYENTE")))', align=CENTER)
    put(RK, r, 14, f'=IF(OR(K{r}="",K{r}="NC"),"",RANK(K{r},$K$4:$K$43))', align=CENTER)
for lo, hi, nom, col, fillc in M.BANDAS:
    RK.conditional_formatting.add("G4:K43", CellIsRule(operator="between", formula=[str(lo), str(hi)], fill=PatternFill("solid", fgColor=fillc), font=Font(name=FONT, size=10, bold=True, color=col)))
for i, (emp, fc) in enumerate((("EJEMPLO_A", dt.date(2026, 6, 30)), ("EJEMPLO_B", dt.date(2026, 6, 30))), 4):
    RK.cell(row=i, column=1, value=emp).fill = FILL_EX
    RK.cell(row=i, column=2, value=fc).fill = FILL_EX
    RK.cell(row=i, column=1).font = Font(name=FONT, size=10, italic=True, color="616161")
put(RK, 45, 1, "Nota: las filas EJEMPLO_A y EJEMPLO_B son ficticias y existen para mostrar las fórmulas. Borrarlas junto con sus filas de 14_HISTORICO antes de usar.", font=F_NOTE, border=False)
RK.merge_cells("A45:N45")

# =====================================================================================
# 17_JURISDICCIONES
# =====================================================================================
J = sheet("17_JURISDICCIONES", [20, 20, 34, 24, 44, 36, 14, 50], freeze="A4")
title(J, "17_JURISDICCIONES — Equivalencia de fuentes N1 y N3 por país", "Si un país no figura, aplicar la regla general del documento 04 y añadir la fila. El tope de confianza FIN se aplica salvo cuentas auditadas verificadas.")
header(J, 3, ["PAIS", "NORMA_CONTABLE_HABITUAL", "REGISTRO MERCANTIL / CUENTAS (N1)", "REGULADOR DE MERCADOS (N1)", "CONTRATACION PUBLICA Y DEFENSA (N1)", "INSTITUCIONAL (N3) Y SECTORIAL (N4)", "TOPE_CONF_FIN", "NOTA"])
for i, row in enumerate(M.JURISDICCIONES, 4):
    for c, v in enumerate(row, 1):
        put(J, i, c, v, font=F_BOLD if c == 1 else F_BASE)
    J.row_dimensions[i].height = 44
J.auto_filter.ref = f"A3:H{3 + len(M.JURISDICCIONES)}"

# =====================================================================================
# ESTÁNDAR VISUAL (20–24)
# =====================================================================================
LP = sheet("20_LEER_PRIMERO", [34, 110], freeze=None)
title(LP, f"DCIP {M.VERSION} — ESTÁNDAR VISUAL v{M.VERSION_ARCH} · LEER PRIMERO", "La arquitectura gráfica es cerrada y no depende del volumen de contenido.")
lp = [
    ("REGLA CENTRAL", "CONTENIDO VARIABLE / ARQUITECTURA INVARIABLE."),
    ("FUENTE DE VERDAD DEL CONTENIDO", "04_RESULTADO, 07_FICHA, 08–11 y 12_FUENTES de este archivo, con 06_AUDITORIA en VÁLIDA."),
    ("FUENTE DE VERDAD DE LA FORMA", "21_PLANTILLA_VISUAL + 22_REGLAS_LEGIBILIDAD + 24_MAPA_M18."),
    ("SI NO CABE EL TEXTO", "Condensar y sintetizar dentro de la caja. Nunca cambiar 16:9, retícula, módulos o posiciones."),
    ("SI FALTA UN DATO", "Mantener el módulo y usar N/D, NO DISPONIBLE o EVIDENCIA INSUFICIENTE."),
    ("SI UNA DIMENSION ES NO CONCLUYENTE", "M14 muestra 'NC' en esa dimensión y 'NO CONCLUYENTE' en el DCIP global. No se rellena con estimaciones."),
    ("SI HAY RANGO", "M14 muestra 'nota (mín–máx)' y la categoría como rango, tal como salen de 04_RESULTADO."),
    ("SI FALLA UN CONTROL CRÍTICO", "La infografía es NO VÁLIDA y debe regenerarse tras corregir el expediente."),
    ("COLOR", "Dimensiones: FIN verde, IND azul, DEF naranja, PRO morado, solo para identificar módulos. Bandas DCIP: escala neutra de grises y navy; nunca los colores de dimensión."),
    ("ORDEN DE USO", "1) expediente validado → 2) plantilla → 3) maquetación → 4) control final → 5) versión final."),
]
for i, (a, b) in enumerate(lp, 4):
    put(LP, i, 1, a, font=F_BOLD)
    put(LP, i, 2, b)

PV = sheet("21_PLANTILLA_VISUAL", [18] * 12, freeze=None)
PV["A1"] = f"DCIP {M.VERSION}  |  PLANTILLA VISUAL MAESTRA v{M.VERSION_ARCH}  (16:9, 12 columnas, M01–M18)"
PV["A1"].font = F_TITLE
PV.merge_cells("A1:L1")


def modulo(ws, row, c1, c2, code, name, body, dim=None, height=70):
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    cell = ws.cell(row=row, column=c1, value=f"{code}  |  {name}")
    cell.font = F_HDR
    cell.fill = PatternFill("solid", fgColor=M.DIMENSIONES[dim]["color"] if dim else M.NAVY)
    cell.alignment = CENTER
    ws.merge_cells(start_row=row + 1, start_column=c1, end_row=row + 1, end_column=c2)
    b = ws.cell(row=row + 1, column=c1, value=body)
    b.font = F_BASE
    b.alignment = WRAP
    b.border = BORDER
    ws.row_dimensions[row + 1].height = height


# Cabecera
PV.merge_cells("A3:L3")
PV["A3"] = "EMPRESA [NOMBRE]  ·  PAÍS  ·  FECHA DE CORTE [DD/MM/AAAA]  ·  PERÍODO  ·  MONEDA / NORMA CONTABLE  ·  PERÍMETRO PRINCIPAL  ·  PERÍMETRO SECUNDARIO [SI PROCEDE]"
PV["A3"].font = F_HDR; PV["A3"].fill = FILL_HDR; PV["A3"].alignment = CENTER
modulo(PV, 5, 1, 4, "M01", "PERÍMETRO Y DATOS CLAVE", "Empresa · país · jurisdicción · fecha de corte · período · moneda y norma contable · perímetros · peso de defensa (HECHO/ESTIMACIÓN).")
modulo(PV, 5, 5, 8, "M02", "RESUMEN EJECUTIVO", "Síntesis ejecutiva: situación, fortalezas, vulnerabilidades y conclusión DCIP.")
modulo(PV, 5, 9, 12, "M03", "SITUACIÓN EMPRESARIAL", "Modelo de negocio · escala · presencia · plantilla · posición corporativa.")
modulo(PV, 8, 1, 3, "M04", "PROPIEDAD Y CONTROL", "Propiedad · accionistas · control efectivo · gobierno corporativo.")
modulo(PV, 8, 4, 6, "M05", "FINANZAS", "Ventas · EBIT · caja · deuda neta · liquidez · patrimonio · contratación y cartera firme. Importes en M€ con moneda original entre paréntesis.", "FIN")
modulo(PV, 8, 7, 9, "M06", "DEFENSA Y PROGRAMAS", "Programas · contratos · clientes gubernamentales · sostenimiento · tecnologías críticas · control de exportaciones.", "DEF")
modulo(PV, 8, 10, 12, "M07", "INDUSTRIA Y TECNOLOGÍA", "Instalaciones · plantilla · tecnología · I+D+i · integración.", "IND")
modulo(PV, 11, 1, 3, "M08", "CAPACIDAD INDUSTRIAL", "Instalada → efectiva → demostrada → disponible · cuellos de botella · escalabilidad.", "IND")
modulo(PV, 11, 4, 6, "M09", "MERCADO E INTERNACIONALIZACIÓN", "Mercados · exportaciones · clientes · concentración.")
modulo(PV, 11, 7, 9, "M10", "RIESGOS", "Riesgos materiales con severidad (08_RIESGOS).")
modulo(PV, 11, 10, 12, "M11", "OPORTUNIDADES", "Oportunidades con severidad (09_OPORTUNIDADES).")
modulo(PV, 14, 1, 7, "M12", "PROSPECTIVA", "12–24 meses: BASE / ALCISTA / BAJISTA\n3–5 años: BASE / ALCISTA / BAJISTA\n5–10 años: BASE / ALCISTA / BAJISTA\nImpulsores · riesgos · señales tempranas · factores de invalidación.", "PRO", 90)
modulo(PV, 14, 8, 12, "M13", "ALERTAS TEMPRANAS", "Señales activas de 11_ALERTA_TEMPRANA con trigger y fuente de seguimiento.", None, 90)
modulo(PV, 17, 1, 4, "M14", "SCORING DCIP", "FIN [__] (mín __)   IND [__] (mín __)   DEF [__] (mín __)   PRO [__] (mín __)\nCobertura: FIN __ %  IND __ %  DEF __ %  PRO __ %\nDCIP GLOBAL [__] (rango __–__)   Categoría: [____]\nConfianza: [____]   Ponderación 25/25/25/25", None, 100)
modulo(PV, 17, 5, 6, "M15", "ESTADO EMPRESARIAL", "[CRÍTICO / FRÁGIL / ESTABLE / SÓLIDO / MUY SÓLIDO]\nSíntesis breve, separada del juicio estratégico.", None, 100)
modulo(PV, 17, 7, 8, "M16", "ATRACTIVIDAD ESTRATÉGICA", "[BAJA / MODERADA / RELEVANTE / ALTA / MUY ALTA]\nRelevancia, posición y perspectivas.", None, 100)
modulo(PV, 17, 9, 10, "M17", "JUICIO DCIP", "Conclusión integrada, diferenciando hechos, estimaciones, inferencias y expectativas.", None, 100)
modulo(PV, 17, 11, 12, "M18", "FUENTES", "Fuentes principales (FUENTE_ID, nivel, fecha) y trazabilidad. Prioridad N1–N3.", None, 100)
PV.merge_cells("A20:L20")
PV["A20"] = f"DCIP {M.VERSION}  |  LA CALIDAD DE LA CONCLUSIÓN NUNCA PUEDE SUPERAR LA CALIDAD DE LA EVIDENCIA."
PV["A20"].font = F_HDR; PV["A20"].fill = FILL_HDR; PV["A20"].alignment = CENTER

RL = sheet("22_REGLAS_LEGIBILIDAD", [22, 90, 14], freeze=None)
title(RL, f"DCIP {M.VERSION} — REGLAS VISUALES v{M.VERSION_ARCH}")
header(RL, 3, ["Regla", "Especificación", "Carácter"])
reglas = [
    ("Formato", "Horizontal 16:9 obligatorio."),
    ("Retícula", "12 columnas fijas; posiciones M01–M18 invariables (24_MAPA_M18)."),
    ("Arquitectura", "Exactamente M01–M18; no eliminar, fusionar, añadir ni reordenar."),
    ("Legibilidad", "Cuerpo mínimo 10 pt; preferente 11–13 pt."),
    ("Densidad", "Si no cabe: condensar, eliminar redundancias y abreviar. No reducir indefinidamente la tipografía."),
    ("Color de dimensión", "FIN verde; IND azul; DEF naranja; PRO morado. Solo para identificar módulos. Navy para estructura."),
    ("Color de banda", "Crítica, Débil, Estable, Favorable, Muy favorable en escala neutra (rojo oscuro → gris → navy). Nunca los colores de dimensión."),
    ("Contenido", "Solo información del expediente validado. No investigar durante el renderizado."),
    ("Datos ausentes", "N/D, no disponible o evidencia insuficiente. Nunca inventar."),
    ("Scoring", "Copiar exactamente de 04_RESULTADO: nota, mínimo, cobertura, categoría y confianza. No recalcular. Mostrar NC y rangos tal cual."),
    ("Unidades", "M€ con moneda original entre paréntesis cuando no sea euro; %; FTE. Punto de miles y coma decimal."),
    ("Estado / atractividad", "M15 y M16 independientes y separados, con su escala de cinco niveles."),
    ("Prospectiva", "M12 conserva 12–24 meses, 3–5 años y 5–10 años; BASE / ALCISTA / BAJISTA."),
    ("Fuentes", "M18 siempre visible y legible, con FUENTE_ID y nivel."),
    ("Estética", "Sobria, técnica y profesional, propia de inteligencia empresarial y defensa."),
    ("Control final", "Comprobar módulos, cifras, fechas, nombres, scoring, rangos, cobertura, confianza, fuentes, ortografía y ausencia de invenciones."),
]
for i, (a, b) in enumerate(reglas, 4):
    put(RL, i, 1, a, font=F_BOLD)
    put(RL, i, 2, b)
    put(RL, i, 3, "OBLIGATORIO", align=CENTER)

PG = sheet("23_PROMPT_GENERACION", [120], freeze=None)
title(PG, f"PROMPT ÚNICO — INFOGRAFÍA DCIP {M.VERSION}")
prompt = f"""INFOGRAFÍA DCIP {M.VERSION} — [EMPRESA]

Usa exclusivamente:
1) el expediente DCIP {M.VERSION} validado de [EMPRESA] (hojas 04_RESULTADO, 07_FICHA, 08_RIESGOS, 09_OPORTUNIDADES, 10_PROSPECTIVA, 11_ALERTA_TEMPRANA y 12_FUENTES, con 06_AUDITORIA en VÁLIDA) como fuente de verdad del contenido;
2) la PLANTILLA VISUAL MAESTRA v{M.VERSION_ARCH} como fuente de verdad de arquitectura y diseño.

CONTENIDO VARIABLE / ARQUITECTURA INVARIABLE.

Aplicar exactamente: horizontal 16:9; retícula de 12 columnas; M01–M18 en sus posiciones; ningún módulo eliminado, fusionado, dividido, añadido, movido o reordenado; M15 Estado empresarial separado de M16 Atractividad estratégica; M17 Juicio DCIP y M18 Fuentes siempre presentes; FIN verde, IND azul, DEF naranja, PRO morado solo como identificadores de módulo; bandas DCIP en escala neutra.

LEGIBILIDAD: cuerpo mínimo 10 pt, preferente 11–13 pt. Si no cabe, condensar y abreviar sin cambiar significado. Nunca solucionar falta de espacio con texto microscópico ni cambiando la arquitectura.

CONTENIDO: copiar exactamente cifras, fechas, nombres, programas, contratos, puntuaciones, mínimos, coberturas, categorías, confianza y fuentes del expediente. Mostrar 'NC' y 'NO CONCLUYENTE' donde el expediente lo indique y los rangos tal como aparecen. No volver a analizar, no recalcular scoring, no añadir información nueva y no inventar datos, palabras, siglas, programas, contratos, fechas o conclusiones. Mantener las distinciones entre ingresos, contratación, cartera, cartera firme, pipeline, contrato, opción, MoU, hecho, estimación, inferencia, expectativa e hipótesis. Importes en M€ con moneda original entre paréntesis cuando no sea euro. Si falta evidencia: N/D o evidencia insuficiente.

CONTROL FINAL: 16:9 · 12 columnas · M01–M18 · posiciones · legibilidad · scoring con rangos y cobertura · confianza · fechas · cifras · nombres · Estado · Atractividad · Juicio · Fuentes · ortografía · ausencia de invenciones. Si alguna condición falla, regenerar antes de entregar."""
put(PG, 3, 1, prompt, border=False)
PG.row_dimensions[3].height = 420

MP = sheet("24_MAPA_M18", [10, 34, 8, 12, 60], freeze="A2")
header(MP, 1, ["Módulo", "Denominación", "Zona", "Columnas", "Contenido"])
for i, (a, b, c, d, e) in enumerate(M.MODULOS_VISUALES, 2):
    put(MP, i, 1, a, font=F_BOLD, align=CENTER)
    put(MP, i, 2, b)
    put(MP, i, 3, c, align=CENTER)
    put(MP, i, 4, d, align=CENTER)
    put(MP, i, 5, e)

# Orden final de hojas
order = ["00_README", "01_CONFIG", "02_INDICADORES", "03_SCORING", "04_RESULTADO", "05_EVIDENCIA", "06_AUDITORIA", "07_FICHA", "08_RIESGOS",
         "09_OPORTUNIDADES", "10_PROSPECTIVA", "11_ALERTA_TEMPRANA", "12_FUENTES", "13_CAMBIOS", "14_HISTORICO", "15_RANKING", "16_LISTAS",
         "17_JURISDICCIONES", "20_LEER_PRIMERO", "21_PLANTILLA_VISUAL", "22_REGLAS_LEGIBILIDAD", "23_PROMPT_GENERACION", "24_MAPA_M18"]
wb._sheets = [wb[n] for n in order]
for ws in wb.worksheets:
    ws.sheet_properties.tabColor = M.NAVY if not ws.title.startswith("2") else "78909C"
wb["02_INDICADORES"].sheet_properties.tabColor = "455A64"
wb["03_SCORING"].sheet_properties.tabColor = "FFC107"
wb["04_RESULTADO"].sheet_properties.tabColor = "FFC107"
wb.save(OUT)
print("OK", OUT)
