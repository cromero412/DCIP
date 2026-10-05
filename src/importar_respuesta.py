# -*- coding: utf-8 -*-
"""
Vuelca la respuesta de un asistente de IA (formato de la sección 7 del prompt operativo) en una copia de la
arquitectura maestra DCIP 2.3. Solo escribe en celdas de entrada; las fórmulas quedan intactas.

Uso:
    python3 importar_respuesta.py respuesta.md DCIP_2.3_Arquitectura_Maestra_v3.0.xlsx expediente_salida.xlsx

Después, abrir el archivo en Excel (recalcula al abrir) o pasar recalc.py de LibreOffice.
"""
import datetime as dt
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

CONFIG_ROWS = {"EMPRESA": 4, "PAIS": 5, "JURISDICCION": 6, "FECHA_CORTE": 7, "PERIODO_FINANCIERO": 8, "CIERRE_FISCAL": 9, "NORMA_CONTABLE": 10,
               "MONEDA_ORIGINAL": 11, "TIPO_CAMBIO_EUR": 12, "FUENTE_TIPO_CAMBIO": 13, "PERIMETRO_PRINCIPAL": 14, "PERIMETRO_SECUNDARIO": 15,
               "PESO_DEFENSA_INGRESOS": 16, "PESO_DEFENSA_TIPO": 17, "VERSION_EXPEDIENTE": 24, "ANALISTA": 25, "FECHA_ANALISIS": 26}

# columnas de entrada de cada hoja (nombre de columna en la respuesta → índice de columna en la hoja)
SCORING_COLS = {"PUNTUACION": 5, "RANGO_MIN": 6, "RANGO_MAX": 7, "TIPO": 12, "NIVEL_FUENTE": 13, "ANTIG_MESES": 14, "ES_CUENTAS": 15,
                "CONFIANZA": 17, "JUSTIFICACION": 19, "EVIDENCIA_IDS": 20, "CONTRADICCION": 21, "OBSERVACIONES": 22}
EVID_COLS = {"EVIDENCIA_ID": 1, "ID_INDICADOR": 2, "FECHA_EVIDENCIA": 5, "FUENTE_ID": 8, "ES_CUENTAS": 10, "TIPO": 11, "CITA_EXTRACTO": 12,
             "REFERENCIA": 13, "PERIMETRO": 14, "MONEDA": 15, "IDIOMA": 16, "TRADUCCION": 17, "CONTRADICCION": 19, "RESOLUCION": 20, "OBSERVACIONES": 21}
FUENTE_COLS = {"FUENTE_ID": 1, "NIVEL": 2, "JURISDICCION": 3, "ORGANISMO": 4, "TITULO": 5, "FECHA": 6, "REFERENCIA": 7, "IDIOMA": 8, "PERIMETRO": 9,
               "NORMA_CONTABLE": 10, "MONEDA": 11, "AFIRMACIONES_SUSTENTADAS": 12, "OBSERVACIONES": 13}
RIESGO_COLS = {"ID": 1, "RIESGO": 2, "OPORTUNIDAD": 2, "PROBABILIDAD": 3, "IMPACTO": 4, "INDICADOR_ADELANTADO": 6, "CONDICION_HABILITANTE": 6,
               "MITIGACION / CONDICION": 7, "HORIZONTE": 8, "EVIDENCIA_ID": 9, "TIPO": 10, "OBSERVACIONES": 11}
PROSP_COLS = {"DESCRIPCION": 3, "IMPULSORES": 4, "RIESGOS": 5, "SEÑALES_TEMPRANAS": 6, "FACTORES_INVALIDACION": 7, "EVIDENCIA_ID": 8, "TIPO": 9}
ALERTA_COLS = {"PRIORIDAD": 1, "AMBITO": 2, "INDICADOR / SEÑAL": 3, "TRIGGER": 4, "FUENTE DE SEGUIMIENTO": 5, "FRECUENCIA": 6, "ACCION": 7, "ESTADO": 8}
AUDIT_COLS = {"ESTADO": 5, "EVIDENCIA": 6, "OBSERVACIONES": 7}
NUMERIC = {"PUNTUACION", "RANGO_MIN", "RANGO_MAX", "ANTIG_MESES", "TIPO_CAMBIO_EUR", "PESO_DEFENSA_INGRESOS", "VALOR_MONEDA_ORIGINAL_M"}
DATES = {"FECHA_EVIDENCIA", "FECHA", "FECHA_CORTE", "FECHA_ANALISIS"}


def parse_tables(text):
    """Devuelve {seccion: [dict(fila)]} a partir de '## SECCION' seguido de una tabla Markdown."""
    out, section, headers = {}, None, None
    for line in text.splitlines():
        s = line.strip()
        m = re.match(r"^##\s+([A-ZÁÉÍÓÚÑ_ ]+)\s*$", s)
        if m:
            section, headers = m.group(1).strip(), None
            out.setdefault(section, [])
            continue
        if section is None or not s.startswith("|"):
            continue
        cells = [c.strip() for c in s.strip("|").split("|")]
        if headers is None:
            headers = cells
            continue
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
            continue
        row = {h: (v if v not in ("", "...", "N/D", "—") else None) for h, v in zip(headers, cells)}
        if any(v is not None for v in row.values()):
            out[section].append(row)
    return out


def conv(key, v):
    if v is None:
        return None
    if key in NUMERIC:
        try:
            return float(v.replace(",", "."))
        except ValueError:
            return v
    if key in DATES:
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return dt.datetime.strptime(v, fmt).date()
            except ValueError:
                pass
    return v


def write_rows(ws, rows, colmap, start_row=2):
    for i, row in enumerate(rows):
        for key, col in colmap.items():
            if key in row and row[key] is not None:
                ws.cell(row=start_row + i, column=col, value=conv(key, row[key]))


def main(resp_path, template_path, out_path):
    text = Path(resp_path).read_text(encoding="utf-8")
    T = parse_tables(text)
    wb = load_workbook(template_path)
    log = []

    # CONFIG
    C = wb["01_CONFIG"]
    for row in T.get("CONFIG", []):
        k, v = row.get("CAMPO"), row.get("VALOR")
        if k in CONFIG_ROWS and v is not None:
            C.cell(row=CONFIG_ROWS[k], column=2, value=conv(k, v))
    log.append(f"CONFIG: {len(T.get('CONFIG', []))} campos")

    # FUENTES, EVIDENCIA
    write_rows(wb["12_FUENTES"], T.get("FUENTES", []), FUENTE_COLS)
    write_rows(wb["05_EVIDENCIA"], T.get("EVIDENCIA", []), EVID_COLS)
    log.append(f"FUENTES: {len(T.get('FUENTES', []))} filas · EVIDENCIA: {len(T.get('EVIDENCIA', []))} filas")

    # SNAPSHOT (por nombre de magnitud)
    F = wb["07_FICHA"]
    snap_rows = {F.cell(row=r, column=1).value: r for r in range(4, 40) if F.cell(row=r, column=1).value}
    n = 0
    for row in T.get("SNAPSHOT", []):
        r = snap_rows.get(row.get("MAGNITUD"))
        if r:
            if row.get("VALOR_MONEDA_ORIGINAL_M") is not None:
                F.cell(row=r, column=2, value=conv("VALOR_MONEDA_ORIGINAL_M", row["VALOR_MONEDA_ORIGINAL_M"]))
            if row.get("EVIDENCIA_ID"):
                F.cell(row=r, column=4, value=row["EVIDENCIA_ID"])
            if row.get("TIPO"):
                F.cell(row=r, column=5, value=row["TIPO"])
            n += 1
    log.append(f"SNAPSHOT: {n} magnitudes")

    # JUICIOS (texto de la ficha + estado/atractividad en 04_RESULTADO)
    ficha_rows = {F.cell(row=r, column=1).value: r for r in range(1, F.max_row + 1) if F.cell(row=r, column=1).value}
    R = wb["04_RESULTADO"]
    for row in T.get("JUICIOS", []):
        k, v = row.get("CAMPO"), row.get("CONTENIDO")
        if v is None:
            continue
        if k in ficha_rows:
            F.cell(row=ficha_rows[k], column=2, value=v)
        if k == "ESTADO_EMPRESARIAL":
            R["B18"] = v.split(".")[0].strip().upper() if v else None
        if k == "ATRACTIVIDAD_ESTRATEGICA":
            R["B19"] = v.split(".")[0].strip().upper() if v else None

    # SCORING (por ID)
    S = wb["03_SCORING"]
    import dcip_model as M
    id_row = {ind["id"]: r for r, ind in enumerate(M.INDICADORES, 2)}
    n = 0
    for row in T.get("SCORING", []):
        r = id_row.get(row.get("ID"))
        if not r:
            log.append(f"  AVISO: ID desconocido en SCORING: {row.get('ID')}")
            continue
        for key, col in SCORING_COLS.items():
            if row.get(key) is not None:
                S.cell(row=r, column=col, value=conv(key, row[key]))
        n += 1
    log.append(f"SCORING: {n} indicadores")

    # RIESGOS, OPORTUNIDADES, ALERTAS, PROSPECTIVA, AUDITORIA
    write_rows(wb["08_RIESGOS"], T.get("RIESGOS", []), RIESGO_COLS)
    write_rows(wb["09_OPORTUNIDADES"], T.get("OPORTUNIDADES", []), RIESGO_COLS)
    write_rows(wb["11_ALERTA_TEMPRANA"], T.get("ALERTAS", []), ALERTA_COLS)
    P = wb["10_PROSPECTIVA"]
    prow = {(P.cell(row=r, column=1).value, P.cell(row=r, column=2).value): r for r in range(2, 11)}
    for row in T.get("PROSPECTIVA", []):
        r = prow.get((row.get("ESCENARIO"), row.get("HORIZONTE")))
        if r:
            for key, col in PROSP_COLS.items():
                if row.get(key) is not None:
                    P.cell(row=r, column=col, value=row[key])
    A = wb["06_AUDITORIA"]
    for row in T.get("AUDITORIA", []):
        try:
            cid = int(float(row.get("CONTROL_ID")))
        except (TypeError, ValueError):
            continue
        r = 6 + cid
        for key, col in AUDIT_COLS.items():
            if row.get(key) is not None:
                A.cell(row=r, column=col, value=row[key])
    log.append(f"RIESGOS: {len(T.get('RIESGOS', []))} · OPORTUNIDADES: {len(T.get('OPORTUNIDADES', []))} · ALERTAS: {len(T.get('ALERTAS', []))} · PROSPECTIVA: {len(T.get('PROSPECTIVA', []))} · AUDITORIA: {len(T.get('AUDITORIA', []))}")

    # Comprobaciones de integridad de la respuesta
    fuentes = {r.get("FUENTE_ID") for r in T.get("FUENTES", [])}
    evid = {r.get("EVIDENCIA_ID") for r in T.get("EVIDENCIA", [])}
    for r in T.get("EVIDENCIA", []):
        if r.get("FUENTE_ID") and r["FUENTE_ID"] not in fuentes:
            log.append(f"  AVISO: {r.get('EVIDENCIA_ID')} cita FUENTE_ID inexistente {r['FUENTE_ID']}")
    for r in T.get("SCORING", []):
        for e in re.split(r"[,; ]+", r.get("EVIDENCIA_IDS") or ""):
            if e and e not in evid:
                log.append(f"  AVISO: {r.get('ID')} cita EVIDENCIA_ID inexistente {e}")
    if len(T.get("SCORING", [])) != len(M.INDICADORES):
        log.append(f"  AVISO: SCORING tiene {len(T.get('SCORING', []))} filas; se esperaban {len(M.INDICADORES)}")

    wb.save(out_path)
    print("\n".join(log))
    print("OK", out_path, "(abrir en Excel para recalcular, o pasar recalc.py)")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print(__doc__)
        sys.exit(1)
    sys.path.insert(0, str(Path(__file__).parent))
    main(*sys.argv[1:])
