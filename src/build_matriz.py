# -*- coding: utf-8 -*-
"""Genera 02_DCIP_2.3_Matriz_de_scoring_v3.0.xlsx a partir de dcip_model.py."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import dcip_model as M

OUT = "/home/user/defensa/DCIP_2.3/02_DCIP_2.3_Matriz_de_scoring_v3.0.xlsx"
FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True, color=M.NAVY)
F_NOTE = Font(name=FONT, size=9, italic=True, color="616161")
FILL_HDR = PatternFill("solid", fgColor=M.NAVY)
FILL_GREY = PatternFill("solid", fgColor="ECEFF1")
THIN = Side(style="thin", color="B0BEC5")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()
wb.remove(wb.active)


def sheet(name, widths, freeze="A2"):
    ws = wb.create_sheet(name)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    if freeze:
        ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    return ws


def header(ws, row, labels):
    for c, lab in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=lab)
        cell.font, cell.fill, cell.alignment, cell.border = F_HDR, FILL_HDR, CENTER, BORDER


def put(ws, r, c, v, font=F_BASE, fill=None, align=WRAP, fmt=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font, cell.alignment, cell.border = font, align, BORDER
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    return cell


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_NOTE


N = len(M.INDICADORES)
LAST = N + 1

# ---------------- Matriz ----------------
ws = sheet("Matriz_v3", [6, 6, 30, 7, 16, 44, 30, 30, 30, 30, 30, 10, 16, 46])
header(ws, 1, ["ID", "DIM", "INDICADOR", "PESO", "POLARIDAD", "METRICA", "NIVEL 0", "NIVEL 25", "NIVEL 50", "NIVEL 75", "NIVEL 100", "FUENTE_PREF", "NO_DOBLE_CONTEO", "REGLA DE USO"])
for r, ind in enumerate(M.INDICADORES, 2):
    f = PatternFill("solid", fgColor=M.DIMENSIONES[ind["dim"]]["color_suave"])
    put(ws, r, 1, ind["id"], font=F_BOLD, fill=f)
    put(ws, r, 2, ind["dim"], font=F_BOLD, fill=f)
    put(ws, r, 3, ind["nombre"], font=F_BOLD)
    put(ws, r, 4, ind["peso"], align=CENTER)
    put(ws, r, 5, ind["polaridad"])
    put(ws, r, 6, ind["metrica"])
    for k, lvl in enumerate((0, 25, 50, 75, 100)):
        put(ws, r, 7 + k, ind["rubrica"][lvl])
    put(ws, r, 12, ind["fuente"], align=CENTER)
    put(ws, r, 13, ", ".join(ind["no_doble"]))
    regla = ("Puntuar 0–100 eligiendo el nivel de la rúbrica cuya descripción se cumple con evidencia citada (EVIDENCIA_ID). "
             "Valores intermedios solo si la evidencia sitúa a la empresa entre dos niveles; conservar el rango. Sin evidencia suficiente: no puntuar.")
    if ind["nota"]:
        regla += " " + ind["nota"]
    put(ws, r, 14, regla, font=F_NOTE)
    ws.row_dimensions[r].height = 96
ws.auto_filter.ref = f"A1:N{LAST}"

# ---------------- Reglas ----------------
R = sheet("Reglas", [26, 120], freeze=None)
title(R, f"DCIP {M.VERSION} — Matriz de scoring v{M.VERSION_ARCH} — {M.FECHA_VERSION}")
reglas = [
    ("Versión", f"DCIP {M.VERSION} Scoring Matrix v{M.VERSION_ARCH} — {M.FECHA_VERSION}. Sustituye a la v2.01 (48 indicadores)."),
    ("Estructura", f"{N} indicadores: FIN 10, IND 10, DEF 12, PRO 9. Cada dimensión suma 100 puntos."),
    ("Pesos de dimensión", "FIN 25 % · IND 25 % · DEF 25 % · PRO 25 %. Fijos."),
    ("Rúbrica", "Cinco niveles por indicador: 0 / 25 / 50 / 75 / 100. El nivel se elige por la descripción que la evidencia acredita. 100 es siempre la situación más favorable, también en los indicadores cuya métrica es 'menor es mejor'."),
    ("Puntuación intermedia", "Permitida solo si la evidencia sitúa a la empresa entre dos niveles; entonces se registra el rango (mínimo y máximo) y el punto central."),
    ("Evidencia", "Cada puntuación cita uno o varios EVIDENCIA_ID. Sin evidencia suficiente el indicador queda vacío: no se imputa."),
    ("Doble conteo", "La columna NO_DOBLE_CONTEO lista los indicadores que no pueden puntuarse con el mismo hecho. Un EVIDENCIA_ID no puede sustentar dos indicadores excluyentes."),
    ("Normalización", "Nota de dimensión = Σ(puntuación × peso) / Σ(pesos con puntuación). Nota prudente = (Σ(puntuación × peso) + 25 × Σ(pesos sin puntuación)) / 100."),
    ("Rango", "Se publica el intervalo [prudente, normalizada]. Si cruza una banda, la categoría se expresa como rango y no se fuerza una banda única."),
    ("Cobertura", "Σ pesos con puntuación / 100. ≥ 80 %: concluyente. 60–79 %: con reservas, confianza máxima MEDIA. < 60 %: dimensión NO CONCLUYENTE y sin DCIP global."),
    ("Redondeo", "Las notas de dimensión y el DCIP se redondean al entero más próximo (mitad hacia arriba) antes de aplicar bandas."),
    ("DCIP", "0,25 FIN + 0,25 IND + 0,25 DEF + 0,25 PRO, sobre las notas normalizadas; límite inferior con las prudentes."),
    ("Bandas", "0–20 Crítica · 21–40 Débil · 41–60 Estable · 61–80 Favorable · 81–100 Muy favorable."),
    ("Confianza", "MUY ALTA / ALTA / MEDIA / BAJA / MUY BAJA. Se deriva del nivel de fuente (N1 → MUY ALTA … N6 → MUY BAJA) y de la antigüedad de la evidencia. Independiente del score."),
    ("Antigüedad", "≤ 12 meses sin penalización; 12–24 meses un nivel menos; > 24 meses dos niveles menos. Para cuentas anuales auditadas los umbrales son 18 y 30 meses."),
    ("Confianza de dimensión", "Media ponderada de la confianza de los indicadores puntuados, redondeada hacia abajo y limitada por la cobertura. Confianza global = la menor de las cuatro."),
    ("Benchmark", "Anclas absolutas en los indicadores con métrica numérica. Donde no es robusto (p. ej. exportación en mercados domésticos muy grandes) comparar con pares homogéneos del mismo país y declarar el ajuste."),
    ("Propiedad pública", "No penaliza. F07 mide apoyo financiero no de mercado, no propiedad."),
    ("Contratación", "No equivale a ingresos."),
    ("Cartera", "No equivale a ingresos futuros garantizados. F08 (cobertura y conversión), D02 (firmeza y composición) y P02 (visibilidad) miden aspectos distintos y no comparten evidencia."),
    ("EBITDA", "Solo publicado o cálculo propio explícitamente reproducible."),
    ("Instrumentos híbridos", "Préstamos participativos y similares no son automáticamente deuda ordinaria."),
    ("Norma contable y moneda", "Ratios en moneda original. Declarar IFRS, US GAAP o GAAP local y su efecto en EBIT, deuda (IFRS 16) y cartera (funded / unfunded)."),
    ("Perímetro de defensa", "Si no hay segmento publicado, estimar el peso de defensa y etiquetarlo ESTIMACIÓN. D02, D03, D07 y D12 se calculan sobre el perímetro de defensa."),
    ("Sanciones y control de exportaciones", "Se evalúan en D10 con listas oficiales (UE, OFAC, ONU, Reino Unido). Su efecto financiero cuantificado va a F10."),
]
header(R, 3, ["REGLA", "ESPECIFICACION"])
for i, (a, b) in enumerate(reglas, 4):
    put(R, i, 1, a, font=F_BOLD)
    put(R, i, 2, b)

# ---------------- Control (fórmulas) ----------------
Cn = sheet("Control", [18, 14, 14, 14], freeze=None)
title(Cn, "Control de integridad de la matriz", "Fórmulas sobre la hoja Matriz_v3. Todas las sumas deben ser 100 y el total de indicadores igual al declarado.")
header(Cn, 3, ["DIMENSION", "SUMA_PESOS", "INDICADORES", "CONTROL"])
for i, d in enumerate(M.DIMENSIONES, 4):
    put(Cn, i, 1, d, font=F_BOLD, fill=PatternFill("solid", fgColor=M.DIMENSIONES[d]["color_suave"]))
    put(Cn, i, 2, f'=SUMIF(Matriz_v3!$B$2:$B${LAST},A{i},Matriz_v3!$D$2:$D${LAST})', align=CENTER)
    put(Cn, i, 3, f'=COUNTIF(Matriz_v3!$B$2:$B${LAST},A{i})', align=CENTER)
    put(Cn, i, 4, f'=IF(B{i}=100,"OK","ERROR")', align=CENTER, font=F_BOLD)
put(Cn, 8, 1, "TOTAL", font=F_BOLD)
put(Cn, 8, 2, "=SUM(B4:B7)", align=CENTER)
put(Cn, 8, 3, "=SUM(C4:C7)", align=CENTER)
put(Cn, 8, 4, f'=IF(AND(B8=400,C8={N}),"OK","ERROR")', align=CENTER, font=F_BOLD)
put(Cn, 10, 1, "IDs únicos", font=F_BOLD)
put(Cn, 10, 2, f'=IF(SUMPRODUCT(1/COUNTIF(Matriz_v3!$A$2:$A${LAST},Matriz_v3!$A$2:$A${LAST}))={N},"OK","DUPLICADOS")', align=CENTER, font=F_BOLD)

# ---------------- Confianza ----------------
Cf = sheet("Confianza", [12, 30, 90, 20], freeze=None)
title(Cf, "Nivel de fuente → confianza máxima; antigüedad → penalización; cobertura → límite")
header(Cf, 3, ["NIVEL", "DENOMINACION", "DESCRIPCION", "CONFIANZA_MAX"])
for i, (cod, nom, desc, cmax) in enumerate(M.NIVELES_FUENTE, 4):
    put(Cf, i, 1, cod, font=F_BOLD, align=CENTER)
    put(Cf, i, 2, nom, font=F_BOLD)
    put(Cf, i, 3, desc)
    put(Cf, i, 4, cmax, align=CENTER)
r0 = 4 + len(M.NIVELES_FUENTE) + 1
header(Cf, r0, ["", "ANTIGÜEDAD RESPECTO A LA FECHA DE CORTE", "", "PENALIZACION (niveles)"])
for i, (txt, pen) in enumerate(M.ANTIGUEDAD, r0 + 1):
    put(Cf, i, 2, txt)
    put(Cf, i, 4, pen, align=CENTER)
r1 = r0 + 1 + len(M.ANTIGUEDAD) + 1
header(Cf, r1, ["", "COBERTURA DE LA DIMENSION", "EFECTO", ""])
for i, (c, e) in enumerate(M.COBERTURA, r1 + 1):
    put(Cf, i, 2, c, font=F_BOLD)
    put(Cf, i, 3, e)

# ---------------- Bandas y escalas ----------------
B = sheet("Bandas_y_escalas", [20, 8, 8, 90], freeze=None)
title(B, "Bandas DCIP y escalas de juicio")
header(B, 3, ["BANDA", "MIN", "MAX", "USO"])
for i, (lo, hi, nom, col, fillc) in enumerate(M.BANDAS, 4):
    put(B, i, 1, nom, font=Font(name=FONT, size=10, bold=True, color=col), fill=PatternFill("solid", fgColor=fillc))
    put(B, i, 2, lo, align=CENTER)
    put(B, i, 3, hi, align=CENTER)
    put(B, i, 4, "Aplicar sobre la nota entera. Si el rango prudente–normalizado cruza esta banda, publicar el rango.")
r0 = 4 + len(M.BANDAS) + 1
header(B, r0, ["ESTADO EMPRESARIAL", "", "", "DESCRIPCION"])
for i, (t, d) in enumerate(M.ESCALA_ESTADO, r0 + 1):
    put(B, i, 1, t, font=F_BOLD)
    put(B, i, 4, d)
r1 = r0 + 1 + len(M.ESCALA_ESTADO) + 1
header(B, r1, ["ATRACTIVIDAD ESTRATEGICA", "", "", "DESCRIPCION"])
for i, (t, d) in enumerate(M.ESCALA_ATRACTIVIDAD, r1 + 1):
    put(B, i, 1, t, font=F_BOLD)
    put(B, i, 4, d)

# ---------------- Mapeo 2.2 → 2.3 ----------------
Mp = sheet("Mapeo_2.2_a_2.3", [8, 32, 10, 10, 10, 70])
header(Mp, 1, ["DIM", "INDICADOR 2.2", "PESO 2.2", "DESTINO 2.3", "PESO 2.3", "CAMBIO"])
for i, (d, n, p, dest, p2, c) in enumerate(M.MAPEO_22_23, 2):
    put(Mp, i, 1, d, font=F_BOLD, fill=PatternFill("solid", fgColor=M.DIMENSIONES[d]["color_suave"]))
    put(Mp, i, 2, n)
    put(Mp, i, 3, p, align=CENTER)
    put(Mp, i, 4, dest, align=CENTER, font=F_BOLD)
    put(Mp, i, 5, p2 if p2 is not None else "", align=CENTER)
    put(Mp, i, 6, c)
last = 1 + len(M.MAPEO_22_23)
put(Mp, last + 2, 2, "Control de pesos 2.2 por dimensión", font=F_BOLD)
for k, d in enumerate(M.DIMENSIONES):
    put(Mp, last + 3 + k, 1, d, font=F_BOLD)
    put(Mp, last + 3 + k, 3, f'=SUMIF($A$2:$A${last},A{last + 3 + k},$C$2:$C${last})', align=CENTER)
    put(Mp, last + 3 + k, 5, f'=SUMIF($A$2:$A${last},A{last + 3 + k},$E$2:$E${last})', align=CENTER)
    put(Mp, last + 3 + k, 6, f'=IF(AND(C{last + 3 + k}=100,E{last + 3 + k}=100),"OK","ERROR")', font=F_BOLD)

wb.save(OUT)
print("OK", OUT)
